import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkinter import font as tkfont
from tkinterdnd2 import TkinterDnD, DND_FILES
import urllib.request
import urllib.parse
import urllib.error
import json
import re
import os
import ssl
import time
import html
from bs4 import BeautifulSoup
import pymupdf as fitz  # PyMuPDF

# MS Office 문서 처리를 위한 라이브러리
import docx
from docx.styles.style import ParagraphStyle
from docx.shared import Pt, Mm                  
from docx.enum.text import WD_ALIGN_PARAGRAPH 
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
import openpyxl
import tempfile

LANGUAGES = {
    '한국어': 'ko', '영어': 'en', '일본어': 'ja',
    '중국어(간체)': 'zh-CN', '중국어(번체)': 'zh-TW',
    '독일어': 'de', '프랑스어': 'fr', '스페인어': 'es', '러시아어': 'ru'
}

ALLOWED_FONT_SIZES =[10, 11, 12, 14, 16, 18, 20]
ALLOWED_CHUNK_SIZES =[300, 400, 500, 600, 700, 800]

ALLOWED_ENGINES =[
    "Google 번역 (기본)", "Google 번역 (대체)", "Google (공식 API)",
    "MS 브라우저 번역", "MS (공식 API)", "MyMemory (무료)",
    "OpenAI (공식 API)", "DeepL (공식 API)"
]

ALLOWED_PAPER_SIZES =[
    "A4 (210x297mm)", 
    "A5 (148x210mm)", 
    "B5 (182x257mm)", 
    "B6 (128x182mm)", 
    "국판 (152x225mm)", 
    "46판 (128x188mm)"
]

# 문서 저장 시 선택할 수 있는 폰트 목록
ALLOWED_DOC_FONTS =["맑은 고딕", "바탕", "돋움", "굴림", "궁서", "Arial", "Times New Roman"]

SETTINGS_FILE = "translator_settings.json"

class OverloadException(Exception):
    pass

def load_settings():
    default_settings = {
        "target_lang": "한국어", "font_size": 11, "chunk_size": 400, 
        "engine": "MS 브라우저 번역", "paper_size": "A4 (210x297mm)",
        "openai_api_key": "", "deepl_api_key": "",
        "google_api_key": "", "ms_api_key": "",
        "doc_font_family": "맑은 고딕",
        "doc_font_size": 11,
        "doc_line_spacing": 1.4,
        "doc_para_spacing": 10,
        "doc_margin_x": 15,
        "doc_margin_y": 20,
        "doc_page_numbers": True,
        "last_url": ""  
    }
    
    if os.path.exists(SETTINGS_FILE):
        try:
            with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
                saved = json.load(f)
                
            for key in default_settings:
                if key not in saved:
                    saved[key] = default_settings[key]
                    
            if saved["engine"] not in ALLOWED_ENGINES:
                saved["engine"] = default_settings["engine"]
                
            current_paper = str(saved.get("paper_size", "")).strip()
            if current_paper not in ALLOWED_PAPER_SIZES:
                mapped_paper = default_settings["paper_size"]
                for allowed in ALLOWED_PAPER_SIZES:
                    if current_paper and current_paper in allowed:
                        mapped_paper = allowed
                        break
                saved["paper_size"] = mapped_paper

            return saved
        except Exception:
            return default_settings
            
    return default_settings

def save_current_config():
    settings = {
        "target_lang": lang_combobox.get() if 'lang_combobox' in globals() else current_settings.get("target_lang", "한국어"), 
        "font_size": font_size_var.get() if 'font_size_var' in globals() else current_settings.get("font_size", 11),
        "chunk_size": chunk_size_var.get() if 'chunk_size_var' in globals() else current_settings.get("chunk_size", 400),
        "engine": engine_var.get() if 'engine_var' in globals() else current_settings.get("engine", "MS 브라우저 번역"),
        "paper_size": paper_size_var.get() if 'paper_size_var' in globals() else current_settings.get("paper_size", "A4 (210x297mm)"),
        "openai_api_key": current_settings.get("openai_api_key", ""),
        "deepl_api_key": current_settings.get("deepl_api_key", ""),
        "google_api_key": current_settings.get("google_api_key", ""),
        "ms_api_key": current_settings.get("ms_api_key", ""),
        "doc_font_family": doc_font_family_var.get() if 'doc_font_family_var' in globals() else current_settings.get("doc_font_family", "맑은 고딕"),
        "doc_font_size": doc_font_size_var.get() if 'doc_font_size_var' in globals() else current_settings.get("doc_font_size", 11),
        "doc_line_spacing": doc_line_spacing_var.get() if 'doc_line_spacing_var' in globals() else current_settings.get("doc_line_spacing", 1.4),
        "doc_para_spacing": doc_para_spacing_var.get() if 'doc_para_spacing_var' in globals() else current_settings.get("doc_para_spacing", 10),
        "doc_margin_x": doc_margin_x_var.get() if 'doc_margin_x_var' in globals() else current_settings.get("doc_margin_x", 15),
        "doc_margin_y": doc_margin_y_var.get() if 'doc_margin_y_var' in globals() else current_settings.get("doc_margin_y", 20),
        "doc_page_numbers": doc_page_numbers_var.get() if 'doc_page_numbers_var' in globals() else current_settings.get("doc_page_numbers", True),
        "last_url": current_settings.get("last_url", "")
    }
    try:
        with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump(settings, f, ensure_ascii=False, indent=4)
    except Exception:
        pass

def save_settings():
    save_current_config()
    root.destroy()

# ----------------- 파일 저장 로직 -----------------
def save_to_file(content, default_name):
    if not content.strip():
        messagebox.showwarning("알림", "저장할 내용이 없습니다.")
        return
        
    filetypes =[
        ("텍스트 문서", "*.txt"),
        ("EPUB 전자책", "*.epub"),
        ("PDF 문서", "*.pdf"),
        ("Word 문서", "*.docx"),
        ("PowerPoint 문서", "*.pptx"),
        ("모든 파일", "*.*")
    ]
    file_path = filedialog.asksaveasfilename(defaultextension=".txt", initialfile=default_name, filetypes=filetypes)
    if not file_path:
        return

    try:
        ext = os.path.splitext(file_path)[1].lower()
        paragraphs =[p.strip() for p in content.split('\n') if p.strip()]
        paper_size = paper_size_var.get()
        
        dimensions = {
            "A4 (210x297mm)": (210, 297), "A5 (148x210mm)": (148, 210), 
            "B5 (182x257mm)": (182, 257), "B6 (128x182mm)": (128, 182), 
            "국판 (152x225mm)": (152, 225), "46판 (128x188mm)": (128, 188)
        }
        width, height = dimensions.get(paper_size, (210, 297))

        # 사용자 설정 불러오기
        d_font_fam = doc_font_family_var.get()
        d_font_sz = doc_font_size_var.get()
        d_line_sp = doc_line_spacing_var.get()
        d_para_sp = doc_para_spacing_var.get()
        d_marg_x = doc_margin_x_var.get()
        d_marg_y = doc_margin_y_var.get()
        d_page_num = doc_page_numbers_var.get()

        if ext == '.docx':
            doc = docx.Document()
            style = doc.styles['Normal']
            if not isinstance(style, ParagraphStyle):
                raise TypeError("문단 스타일이 필요합니다.")
            style.font.name = d_font_fam
            style.font.size = Pt(d_font_sz)
            style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), d_font_fam)
            
            for section in doc.sections:
                section.page_width = Mm(width)
                section.page_height = Mm(height)
                section.left_margin, section.right_margin = Mm(d_marg_x), Mm(d_marg_x)
                section.top_margin, section.bottom_margin = Mm(d_marg_y), Mm(d_marg_y)
                
                if d_page_num:
                    footer = section.footer
                    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = para.add_run()
                    fldChar1 = OxmlElement('w:fldChar')
                    fldChar1.set(qn('w:fldCharType'), 'begin')
                    instrText = OxmlElement('w:instrText')
                    instrText.set(qn('xml:space'), 'preserve')
                    instrText.text = "PAGE"
                    fldChar2 = OxmlElement('w:fldChar')
                    fldChar2.set(qn('w:fldCharType'), 'separate')
                    fldChar3 = OxmlElement('w:fldChar')
                    fldChar3.set(qn('w:fldCharType'), 'end')
                    run._r.append(fldChar1)
                    run._r.append(instrText)
                    run._r.append(fldChar2)
                    run._r.append(fldChar3)

            for text in paragraphs:
                para = doc.add_paragraph(text)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para_format = para.paragraph_format
                para_format.first_line_indent = Pt(d_font_sz * 1.1)
                para_format.line_spacing = d_line_sp
                para_format.space_after = Pt(d_para_sp)
            doc.save(file_path)

        elif ext == '.pdf':
            try:
                from fpdf import FPDF
            except ImportError:
                messagebox.showerror("라이브러리 필요", "PDF 저장을 위해 fpdf2 라이브러리가 필요합니다.\npip install fpdf2")
                return
                
            class CustomPDF(FPDF):
                def footer(self):
                    if d_page_num:
                        self.set_y(-15)
                        self.set_font("helvetica", "I", 9)
                        self.cell(0, 10, str(self.page_no()), align="C")

            pdf = CustomPDF(format=(width, height))
            pdf.set_auto_page_break(auto=False)
            pdf.set_margins(left=d_marg_x, top=d_marg_y, right=d_marg_x)
            pdf.add_page()
            
            font_mapping = {
                "맑은 고딕": "malgun.ttf", "바탕": "batang.ttc", "돋움": "dotum.ttc",
                "굴림": "gulim.ttc", "궁서": "gungsuh.ttc", "Arial": "arial.ttf", "Times New Roman": "times.ttf"
            }
            font_file = font_mapping.get(d_font_fam, "malgun.ttf")
            font_path = f"C:/Windows/Fonts/{font_file}"
            
            if os.path.exists(font_path):
                try:
                    pdf.add_font(d_font_fam, "", font_path)
                    pdf.set_font(d_font_fam, size=d_font_sz)
                except Exception:
                    try:
                        pdf.add_font("MalgunFallback", "", "C:/Windows/Fonts/malgun.ttf")
                        pdf.set_font("MalgunFallback", size=d_font_sz)
                    except:
                        pdf.set_font("helvetica", size=d_font_sz)
            else:
                pdf.set_font("helvetica", size=d_font_sz)
                
            effective_w = width - (d_marg_x * 2)
            line_height = (d_font_sz * 0.35) * d_line_sp
            default_space_w = pdf.get_string_width(' ')
            
            for text in paragraphs:
                text = "    " + text  # 첫 줄 들여쓰기
                
                lines =[]
                cur_line = ""
                for char in text:
                    if pdf.get_string_width(cur_line + char) > effective_w:
                        lines.append(cur_line)
                        if char == ' ':
                            cur_line = "" 
                        else:
                            cur_line = char
                    else:
                        cur_line += char
                if cur_line:
                    lines.append(cur_line)
                
                for i, line in enumerate(lines):
                    if pdf.get_y() + line_height > height - d_marg_y:
                        pdf.add_page()
                    
                    is_last_line = (i == len(lines) - 1)
                    clean_line = line.rstrip() 
                    
                    if is_last_line or not clean_line or clean_line.count(' ') == 0:
                        pdf.set_x(d_marg_x)
                        pdf.cell(effective_w, line_height, clean_line, align='L')
                        pdf.ln(line_height)
                    else:
                        space_count = clean_line.count(' ')
                        text_width = pdf.get_string_width(clean_line)
                        extra_space = (effective_w - text_width) / space_count
                        
                        words = clean_line.split(' ')
                        current_x = d_marg_x
                        
                        for word in words:
                            if word:
                                pdf.set_x(current_x)
                                pdf.cell(pdf.get_string_width(word), line_height, word, align='L')
                            current_x += pdf.get_string_width(word) + default_space_w + extra_space
                        
                        pdf.ln(line_height)
                        
                if pdf.get_y() + (d_para_sp * 0.35) > height - d_marg_y:
                    pdf.add_page()
                else:
                    pdf.ln(d_para_sp * 0.35)
                
            pdf.output(file_path)

        elif ext == '.epub':
            try:
                from ebooklib import epub
            except ImportError:
                messagebox.showerror("라이브러리 필요", "EPUB 저장을 위해 ebooklib가 필요합니다.\npip install EbookLib")
                return
            book = epub.EpubBook()
            book.set_title(default_name)
            book.set_language('ko')
            c1 = epub.EpubHtml(title='본문', file_name='chap_01.xhtml', lang='ko')
            
            html_content = f'<html><head><meta charset="utf-8"></head><body style="margin: {d_marg_y}mm {d_marg_x}mm; font-family: \'{d_font_fam}\', sans-serif; word-break: break-all;">'
            for text in paragraphs:
                html_content += f'<p style="text-indent: 1.2em; text-align: justify; line-height: {d_line_sp}; margin-bottom: {d_para_sp}pt; font-size: {d_font_sz}pt;">{text}</p>'
            html_content += '</body></html>'
            c1.content = html_content
            book.add_item(c1)
            book.toc = [c1]
            book.add_item(epub.EpubNcx())
            book.add_item(epub.EpubNav())
            book.spine =['nav', c1]
            epub.write_epub(file_path, book, {})

        elif ext == '.pptx':
            import pptx.util 
            prs = Presentation()
            original_width = prs.slide_width
            original_height = prs.slide_height
            if original_width is None or original_height is None:
                raise ValueError("슬라이드 크기가 없습니다.")
            slide_width = original_height
            slide_height = original_width
            prs.slide_width = slide_width
            prs.slide_height = slide_height
            blank_layout = prs.slide_layouts[6]
            
            left = pptx.util.Mm(d_marg_x)
            top = pptx.util.Mm(d_marg_y)
            width_ppt = pptx.util.Emu(slide_width - pptx.util.Mm(d_marg_x * 2))
            height_ppt = pptx.util.Emu(slide_height - pptx.util.Mm(d_marg_y * 2))
            if width_ppt <= 0 or height_ppt <= 0:
                raise ValueError("여백이 슬라이드 크기보다 큽니다.")
            
            dynamic_max_chars = int(200000 / (d_font_sz ** 2))
            
            def create_new_slide():
                tf = prs.slides.add_slide(blank_layout).shapes.add_textbox(left, top, width_ppt, height_ppt).text_frame
                tf.word_wrap = True 
                return tf
                
            tf = create_new_slide()
            char_count = 0
            for text in paragraphs:
                eff_len = sum(0.55 if ord(c) < 128 else 1.0 for c in text) + 30
                if char_count + eff_len > dynamic_max_chars and char_count > 0:
                    tf = create_new_slide()
                    char_count = 0
                p_shape = tf.paragraphs[0] if not tf.text and len(tf.paragraphs) > 0 else tf.add_paragraph()
                p_shape.text = text
                p_shape.space_after = pptx.util.Pt(d_para_sp)
                p_shape.line_spacing = d_line_sp
                p_shape.font.size = pptx.util.Pt(d_font_sz)
                p_shape.font.name = d_font_fam  
                char_count += eff_len
            prs.save(file_path)
            
        else:
            with open(file_path, "w", encoding="utf-8") as f:
                f.write("\n\n".join(["    " + p for p in paragraphs]))
                
        status_var.set(f"상태: 파일 저장 완료 ({os.path.basename(file_path)})")
        messagebox.showinfo("저장 완료", "파일이 성공적으로 저장되었습니다.")
    except Exception as e:
        status_var.set("상태: 파일 저장 실패")
        messagebox.showerror("오류", f"저장 중 오류가 발생했습니다.\n{str(e)}")

def save_source_file(): save_to_file(text_input.get("1.0", tk.END), "원문_추출내용")
def save_result_file(): save_to_file(text_output.get("1.0", tk.END), "번역결과")

# ----------------- 번역 엔진 및 API 연동 로직 -----------------
EDGE_TOKEN = None
EDGE_TOKEN_TIME = 0
cancel_requested = False

def translate_openai(text, target_lang_code):
    api_key = current_settings.get("openai_api_key", "").strip()
    if not api_key: return "\n[OpenAI API 키가 설정되지 않았습니다.]\n"
    
    lang_name = "Korean"
    for k, v in LANGUAGES.items():
        if v == target_lang_code: lang_name = k; break

    payload = {
        "model": "gpt-4o-mini", 
        "messages":[
            {"role": "system", "content": f"Translate to {lang_name}. Output ONLY the translation."},
            {"role": "user", "content": text}
        ],
        "temperature": 0.3
    }
    req = urllib.request.Request("https://api.openai.com/v1/chat/completions", data=json.dumps(payload).encode('utf-8'), headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"})
    try:
        with urllib.request.urlopen(req, context=ssl._create_unverified_context()) as response:
            return json.loads(response.read().decode('utf-8'))['choices'][0]['message']['content'].strip()
    except Exception as e: return f"\n[OpenAI 오류: {str(e)}]\n"

def translate_deepl(text, target_lang_code):
    api_key = current_settings.get("deepl_api_key", "").strip()
    if not api_key: return "\n[DeepL API 키가 설정되지 않았습니다.]\n"
    
    dl_map = {'ko': 'KO', 'en': 'EN-US', 'ja': 'JA', 'zh-CN': 'ZH', 'zh-TW': 'ZH'}
    dl_target = dl_map.get(target_lang_code, target_lang_code.upper())
    
    domain = "api-free.deepl.com" if api_key.endswith(":fx") else "api.deepl.com"
    data = urllib.parse.urlencode({'text': text, 'target_lang': dl_target}).encode('utf-8')
    req = urllib.request.Request(f"https://{domain}/v2/translate", data=data, headers={"Authorization": f"DeepL-Auth-Key {api_key}", "Content-Type": "application/x-www-form-urlencoded"})
    try:
        with urllib.request.urlopen(req, context=ssl._create_unverified_context()) as response:
            return json.loads(response.read().decode('utf-8'))['translations'][0]['text']
    except Exception as e: return f"\n[DeepL 오류: {str(e)}]\n"

def translate_google_official(text, target_lang_code):
    api_key = current_settings.get("google_api_key", "").strip()
    if not api_key: return "\n[Google 공식 API 키가 설정되지 않았습니다.]\n"
    
    url = f"https://translation.googleapis.com/language/translate/v2?key={api_key}"
    req = urllib.request.Request(url, data=json.dumps({"q": text, "target": target_lang_code}).encode('utf-8'), headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, context=ssl._create_unverified_context()) as response:
            return html.unescape(json.loads(response.read().decode('utf-8'))['data']['translations'][0]['translatedText'])
    except Exception as e: return f"\n[Google 공식 오류: {str(e)}]\n"

def translate_ms_official(text, target_lang_code):
    api_key_raw = current_settings.get("ms_api_key", "").strip()
    if not api_key_raw: return "\n[MS 공식 API 키가 설정되지 않았습니다.]\n"
        
    parts = api_key_raw.split(',')
    api_key = parts[0].strip()
    region = parts[1].strip() if len(parts) > 1 else ""

    url = f"https://api.cognitive.microsofttranslator.com/translate?api-version=3.0&to={target_lang_code}"
    headers = {"Ocp-Apim-Subscription-Key": api_key, "Content-Type": "application/json"}
    if region: headers["Ocp-Apim-Subscription-Region"] = region
        
    req = urllib.request.Request(url, data=json.dumps([{"Text": text}]).encode('utf-8'), headers=headers)
    try:
        with urllib.request.urlopen(req, context=ssl._create_unverified_context()) as response:
            return json.loads(response.read().decode('utf-8'))[0]['translations'][0]['text']
    except Exception as e: return f"\n[MS 공식 오류: {str(e)}]\n"

def translate_ms_edge(text, target_lang_code):
    global EDGE_TOKEN, EDGE_TOKEN_TIME
    ctx = ssl._create_unverified_context()
    for _ in range(3):
        try:
            if not EDGE_TOKEN or time.time() - EDGE_TOKEN_TIME > 600:
                req_auth = urllib.request.Request("https://edge.microsoft.com/translate/auth", headers={'User-Agent': 'Mozilla/5.0'})
                with urllib.request.urlopen(req_auth, context=ctx) as res:
                    EDGE_TOKEN = res.read().decode('utf-8')
                    EDGE_TOKEN_TIME = time.time()
            
            url = f"https://api-edge.cognitive.microsofttranslator.com/translate?from=&to={target_lang_code}&api-version=3.0"
            req = urllib.request.Request(url, data=json.dumps([{"Text": text}]).encode('utf-8'), headers={'Authorization': f'Bearer {EDGE_TOKEN}', 'Content-Type': 'application/json', 'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, context=ctx) as res:
                return json.loads(res.read().decode('utf-8'))[0]['translations'][0]['text']
        except urllib.error.HTTPError as e:
            if e.code == 429: raise OverloadException()
            elif e.code == 401: EDGE_TOKEN = None; continue
            return f"\n[MS 브라우저 번역 오류 (HTTP {e.code})]\n"
        except OverloadException: raise
        except Exception as e: return f"\n[MS 브라우저 번역 오류: {str(e)}]\n"

def translate_google(text, target_lang_code, client="gtx"):
    url = f"https://translate.googleapis.com/translate_a/single?client={client}&sl=auto&tl={target_lang_code}&dt=t"
    req = urllib.request.Request(url, data=urllib.parse.urlencode({'q': text}).encode('utf-8'), headers={'User-Agent': 'Mozilla/5.0'})
    try:
        with urllib.request.urlopen(req, context=ssl._create_unverified_context()) as res:
            parsed = json.loads(res.read().decode('utf-8'))
            trans_text = "".join([s[0] for s in parsed[0] if s[0]])
            return re.sub(r'[a-f0-9]{32}', '', trans_text, flags=re.IGNORECASE)
    except urllib.error.HTTPError as e:
        if e.code == 429: raise OverloadException()
        return f"\n[Google 번역 오류]\n"
    except OverloadException: raise
    except Exception as e: return f"\n[Google 번역 오류: {str(e)}]\n"

def translate_mymemory(text, target_lang_code):
    url = "https://api.mymemory.translated.net/get"
    data = urllib.parse.urlencode({'q': text, 'langpair': f'Autodetect|{target_lang_code}'}).encode('utf-8')
    req = urllib.request.Request(url, data=data, headers={'User-Agent': 'Mozilla/5.0'})
    try:
        with urllib.request.urlopen(req, context=ssl._create_unverified_context()) as res:
            parsed = json.loads(res.read().decode('utf-8'))
            if parsed.get("responseStatus") == 429: raise OverloadException()
            return parsed.get("responseData", {}).get("translatedText", text)
    except urllib.error.HTTPError as e:
        if e.code == 429: raise OverloadException()
        return f"\n[MyMemory 번역 오류]\n"
    except OverloadException: raise
    except Exception as e: return f"\n[MyMemory 번역 오류: {str(e)}]\n"

def translate_text(text, target_lang_code):
    engine = engine_var.get()
    if engine == "OpenAI (공식 API)": return translate_openai(text, target_lang_code)
    elif engine == "DeepL (공식 API)": return translate_deepl(text, target_lang_code)
    elif engine == "Google (공식 API)": return translate_google_official(text, target_lang_code)
    elif engine == "MS (공식 API)": return translate_ms_official(text, target_lang_code)
    elif engine == "MS 브라우저 번역": return translate_ms_edge(text, target_lang_code)
    elif engine == "MyMemory (무료)": return translate_mymemory(text, target_lang_code)
    elif engine == "Google 번역 (대체)": return translate_google(text, target_lang_code, "dict-chrome-ex")
    else: return translate_google(text, target_lang_code, "gtx")

def stop_translation():
    global cancel_requested
    cancel_requested = True
    status_var.set("상태: 번역 중지 요청됨...")

def on_translate():
    global cancel_requested
    cancel_requested = False

    source_text = text_input.get("1.0", "end-1c")
    if not source_text.strip():
        messagebox.showwarning("알림", "번역할 텍스트를 입력하세요.")
        return
        
    target_lang_code = LANGUAGES.get(lang_combobox.get(), 'ko')
    blocks = []
    current_batch = []
    current_len = 0
    SAFE_LIMIT = chunk_size_var.get()
    
    for line in source_text.split('\n'):
        if current_len + len(line) > SAFE_LIMIT and current_batch:
            blocks.append(current_batch)
            current_batch = [line]
            current_len = len(line)
        else:
            current_batch.append(line)
            current_len += len(line) + 1
            
    if current_batch:
        blocks.append(current_batch)
        
    text_output.delete("1.0", tk.END)

    total_blocks = len(blocks)
    progress_var.set(0)
    progress_bar.config(maximum=total_blocks)

    translate_start_time = time.time()
    speed_text_var.set("0.0블록/s")
    remaining_time_var.set("남은 시간 --:--:--")
    elapsed_time_var.set("소요 시간 00:00:00")

    status_var.set(f"상태:[{engine_var.get()}] {SAFE_LIMIT}자 단위 번역 시작...")
    root.update()

    def update_speed_info(completed):
        elapsed = time.time() - translate_start_time
        elapsed_time_var.set(f"소요 시간 {format_seconds(elapsed)}")

        if elapsed > 0 and completed > 0:
            speed = completed / elapsed
            remaining_blocks = total_blocks - completed
            remaining_seconds = remaining_blocks / speed if speed > 0 else 0

            speed_text_var.set(f"{speed:.2f}블록/s")
            remaining_time_var.set(f"남은 시간 {format_seconds(remaining_seconds)}")
        else:
            speed_text_var.set("0.0블록/s")
            remaining_time_var.set("남은 시간 --:--:--")
    
    try:
        for i, batch in enumerate(blocks):
            if cancel_requested:
                elapsed_time_var.set(f"소요 시간 {format_seconds(time.time() - translate_start_time)}")
                speed_text_var.set("중지됨")
                remaining_time_var.set("남은 시간 --:--:--")
                status_var.set(f"상태: 번역 중지됨 ({i} / {total_blocks})")
                messagebox.showinfo("중지", "번역이 중지되었습니다.")
                return

            status_var.set(f"상태: 번역 진행 중... ({i+1} / {total_blocks}) 블록 처리 중")
            root.update()
            
            block_text = '\n'.join(batch)

            if not block_text.strip():
                text_output.insert(
                    tk.END,
                    '\n' * (len(batch) if i < total_blocks - 1 else len(batch) - 1)
                )

                completed = i + 1
                progress_var.set(completed)
                update_speed_info(completed)

                root.update()
                continue
                
            res_lines = translate_text(block_text, target_lang_code).split('\n')

            if cancel_requested:
                speed_text_var.set("중지됨")
                remaining_time_var.set("남은 시간 --:--:--")
                status_var.set(f"상태: 번역 중지됨 ({i} / {total_blocks})")
                messagebox.showinfo("중지", "번역이 중지되었습니다.")
                return

            if len(res_lines) == len(batch): 
                batch_result = '\n'.join(res_lines)
            else:
                recovered = []
                for line in batch:
                    if cancel_requested:
                        speed_text_var.set("중지됨")
                        remaining_time_var.set("남은 시간 --:--:--")
                        status_var.set(f"상태: 번역 중지됨 ({i} / {total_blocks})")
                        messagebox.showinfo("중지", "번역이 중지되었습니다.")
                        return

                    if not line.strip():
                        recovered.append("")
                    else:
                        recovered.append(translate_text(line, target_lang_code).replace('\n', ' '))
                        time.sleep(0.2)
                        root.update()

                batch_result = '\n'.join(recovered)
                
            text_output.insert(tk.END, batch_result + ('\n' if i < total_blocks - 1 else ''))
            text_input.see(tk.END)
            text_output.see(tk.END)

            completed = i + 1
            progress_var.set(completed)
            update_speed_info(completed)

            status_var.set(f"상태: 번역 진행 중... ({completed} / {total_blocks}) 블록 완료")
            root.update()

            if i < total_blocks - 1:
                time.sleep(0.5)
                
        progress_var.set(total_blocks)
        elapsed_time_var.set(f"소요 시간 {format_seconds(time.time() - translate_start_time)}")
        speed_text_var.set("완료")
        remaining_time_var.set("남은 시간 00:00:00")
        status_var.set("상태: 번역 완료.")

        messagebox.showinfo("완료", f"[{engine_var.get()}] 번역이 완료되었습니다.")
        
    except OverloadException:
        text_output.insert(tk.END, "\n\n[강제 중단] 서버 과부하 감지. 번역을 즉시 중단했습니다.")
        elapsed_time_var.set(f"소요 시간 {format_seconds(time.time() - translate_start_time)}")
        speed_text_var.set("오류")
        remaining_time_var.set("남은 시간 --:--:--")
        status_var.set("상태: 서버 과부하 감지, 번역 즉시 중단됨")
        messagebox.showwarning("강제 중단", "번역 서버 과부하가 감지되었습니다. 잠시 후 시도하세요.")

# ----------------- 파일 / 웹 문서 추출 로직 -----------------
def extract_text_from_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.txt':
        try:
            with open(file_path, "r", encoding="utf-8") as f: return f.read()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="cp949") as f: return f.read()
    elif ext == '.pdf':
        extracted = ""
        doc = fitz.open(file_path)
        for page in doc:
            text = page.get_text("text")
            if isinstance(text, str) and text: extracted += text + "\n\n"
        doc.close()
        pattern = r'[^a-zA-Z0-9\uAC00-\uD7A3\u3131-\u318E\u4E00-\u9FFF\u3040-\u30FF\s\.\,\!\?\'\"\(\)\[\]\-\+\=\/\%\:\;\<\>\~]'
        cleaned = re.sub(pattern, ' ', extracted)
        cleaned = re.sub(r'\n\s*\n', '<P>', cleaned)
        cleaned = cleaned.replace('\n', ' ').replace('<P>', '\n\n')
        return re.sub(r' {2,}', ' ', cleaned)
    elif ext == '.docx': 
        return '\n'.join([p.text for p in docx.Document(file_path).paragraphs if p.text.strip()])
    elif ext == '.pptx': 
        extracted =[]
        for slide in Presentation(file_path).slides:
            for shape in slide.shapes:
                shape_text = getattr(shape, "text", "")
                if isinstance(shape_text, str) and shape_text.strip():
                    extracted.append(shape_text.strip())
            extracted.append("")
        return '\n'.join(extracted)
    elif ext == '.xlsx': 
        extracted =[]
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text =[str(c) for c in row if c is not None]
                if row_text: extracted.append('\t'.join(row_text))
            extracted.append("")
        return '\n'.join(extracted)
    else:
        raise ValueError("지원하지 않는 파일 형식입니다.")

def on_reset(): 
    text_input.delete("1.0", tk.END)
    text_output.delete("1.0", tk.END)
    status_var.set("상태: 초기화되었습니다.")

def load_from_file():
    file_path = filedialog.askopenfilename(filetypes=[("지원 문서", "*.txt;*.pdf;*.docx;*.pptx;*.xlsx"), ("모든 파일", "*.*")])
    if file_path:
        try:
            text = extract_text_from_file(file_path)
            text_input.delete("1.0", tk.END)
            text_input.insert(tk.END, text.strip())
            status_var.set(f"상태: 문서 불러오기 완료 ({os.path.basename(file_path)})")
        except Exception as e: 
            messagebox.showerror("오류", f"문서를 읽지 못했습니다.\n{str(e)}")

def load_from_url():
    result_url =[]
    dialog = tk.Toplevel(root)
    dialog.title("웹사이트 불러오기")
    dialog.geometry("700x200")
    dialog.transient(root)
    dialog.grab_set()
    
    dialog.update_idletasks()
    x = root.winfo_x() + (root.winfo_width() // 2) - 350
    y = root.winfo_y() + (root.winfo_height() // 2) - 90
    dialog.geometry(f"+{x}+{y}")
    
    tk.Label(dialog, text="불러올 웹사이트 주소를 입력하세요:\n(예: www.wikipedia.org)", font=global_font).pack(pady=(20, 10))
    url_entry = tk.Entry(dialog, width=65, font=global_font)
    url_entry.pack(fill=tk.X, padx=30, pady=5)
    
    last_saved_url = current_settings.get("last_url", "")
    if last_saved_url:
        url_entry.insert(0, last_saved_url)
        url_entry.select_range(0, tk.END)
        
    url_entry.focus()
    
    def on_confirm(event=None):
        entered_url = url_entry.get().strip()
        result_url.append(entered_url)
        current_settings["last_url"] = entered_url
        save_current_config()
        dialog.destroy()
        
    url_entry.bind("<Return>", on_confirm)
    
    btn_frame = tk.Frame(dialog)
    btn_frame.pack(pady=15)
    tk.Button(btn_frame, text="확인", font=global_font, command=on_confirm, bg="#4CAF50", fg="white", width=10).pack(side=tk.LEFT, padx=10)
    tk.Button(btn_frame, text="취소", font=global_font, command=dialog.destroy, width=10).pack(side=tk.LEFT, padx=10)
    dialog.wait_window()
    
    if not result_url or not result_url[0]: return
        
    url = result_url[0]
    if not url.startswith(('http://', 'https://')): url = 'https://' + url
        
    status_var.set("상태: 웹사이트를 분석하는 중입니다...")
    root.update()
    
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, context=ssl._create_unverified_context()) as res:
            raw_data = res.read()
            try:
                html_data = raw_data.decode('utf-8')
            except UnicodeDecodeError:
                html_data = raw_data.decode('cp949', errors='ignore')
            
        soup = BeautifulSoup(html_data, 'html.parser')
        
        for element in soup(["script", "style", "nav", "footer", "header", "aside", "noscript", "iframe", "form"]): 
            element.extract()
            
        noise = re.compile(r'(ad|ads|advert|banner|sidebar|menu|popup|comment|share|social|widget|promo)\b', re.IGNORECASE)
        for el in soup.find_all(['div', 'ul', 'li', 'section']):
            classes = el.get('class')
            class_text = classes if isinstance(classes, str) else ' '.join(classes or ())
            element_id = el.get('id')
            id_text = element_id if isinstance(element_id, str) else ' '.join(element_id or ())
            if noise.search(class_text + ' ' + id_text): 
                el.extract()
                
        main_content = soup.find('article') or soup.find('main') or soup.find('body') or soup
        cleaned_text = '\n'.join([line.strip() for line in main_content.get_text(separator='\n').splitlines() if line.strip()])
        
        text_input.delete("1.0", tk.END)
        text_input.insert(tk.END, cleaned_text)
        status_var.set("상태: 웹사이트 텍스트 불러오기 완료.")
    except Exception as e: messagebox.showerror("오류", f"불러오기 실패.\n{str(e)}")

def handle_drop(event):
    files =[f.strip('{}') for f in re.findall(r'\{.*?\}|\S+', event.data)]
    for file_path in files:
        if not os.path.isfile(file_path): continue
        try:
            extracted_text = extract_text_from_file(file_path).strip()
            if not extracted_text: continue
            if text_input.get("1.0", tk.END).strip(): 
                text_input.insert(tk.END, f"\n\n---[추가됨: {os.path.basename(file_path)}] ---\n\n")
            text_input.insert(tk.END, extracted_text)
        except Exception: pass

# ----------------- UI: 커스텀 메뉴 및 다이얼로그 -----------------
class CustomMenuBar:
    def __init__(self, root):
        self.frame = tk.Frame(root, bg="#f0f0f0")
        self.frame.pack(side=tk.TOP, fill=tk.X)
        self.is_active = False
        self.active_menu = None
        self.menus =[]
        root.bind_all("<ButtonPress-1>", self.on_global_click, add="+")
        
    def activate_menu(self, menu):
        if self.active_menu and self.active_menu != menu: self.active_menu.hide()
        self.is_active = True
        self.active_menu = menu
        menu.show()
        
    def deactivate(self):
        if self.active_menu: self.active_menu.hide()
        self.is_active = False
        self.active_menu = None
        
    def on_global_click(self, event):
        if not self.is_active: return
        for m in self.menus:
            if event.widget == m.btn: return
        try:
            top = event.widget.winfo_toplevel()
            if self.active_menu and (top == self.active_menu.top or top in self.active_menu.submenus): return
        except Exception: pass
        self.deactivate()

class CustomDropdownMenu:
    def __init__(self, parent_bar, text, font):
        self.parent_bar = parent_bar
        self.parent_bar.menus.append(self)
        self.font = font
        
        self.btn = tk.Label(parent_bar.frame, text=text, font=font, bg="#f0f0f0", padx=12, pady=4, cursor="hand2")
        self.btn.pack(side=tk.LEFT)
        
        self.top = tk.Toplevel(parent_bar.frame)
        self.top.overrideredirect(True)
        self.top.withdraw()
        self.top.attributes('-topmost', True)
        
        self.frame = tk.Frame(self.top, bg="white", highlightbackground="#cccccc", highlightthickness=1)
        self.frame.pack(fill=tk.BOTH, expand=True)
        self.submenus =[]
        self.items =[]
        self.active_submenu = None
        
        def on_click(e):
            if self.parent_bar.is_active and self.parent_bar.active_menu == self: self.parent_bar.deactivate()
            else: self.parent_bar.activate_menu(self)
        def on_enter(e):
            if self.parent_bar.is_active and self.parent_bar.active_menu != self: self.parent_bar.activate_menu(self)
            elif not self.parent_bar.is_active: self.btn.config(bg="#cce8ff")
        def on_leave(e):
            if not self.parent_bar.is_active or self.parent_bar.active_menu != self: self.btn.config(bg="#f0f0f0")
                
        self.btn.bind("<ButtonPress-1>", on_click)
        self.btn.bind("<Enter>", on_enter)
        self.btn.bind("<Leave>", on_leave)
            
    def show(self):
        self.btn.config(bg="#cce8ff")
        self.top.update_idletasks()
        x = self.btn.winfo_rootx()
        y = self.btn.winfo_rooty() + self.btn.winfo_height()
        self.top.geometry(f"+{x}+{y}")
        self.top.deiconify()
        self.top.lift()
        
    def hide(self):
        self.top.withdraw()
        for sub in self.submenus: sub.withdraw()
        self.active_submenu = None
        self.btn.config(bg="#f0f0f0")
        for item in self.items: item.config(bg="white")

    def add_command(self, label, command, is_special=False):
        color = "blue" if is_special else "black"
        item = tk.Label(self.frame, text=label, font=self.font, bg="white", fg=color, anchor="w", padx=15, pady=8)
        item.pack(fill=tk.X)
        self.items.append(item)
        
        def _enter(e):
            for it in self.items: it.config(bg="white")
            item.config(bg="#cce8ff")
            if self.active_submenu: 
                self.active_submenu.withdraw()
                self.active_submenu = None
                
        item.bind("<Enter>", _enter)
        
        def _click(e):
            self.parent_bar.deactivate()
            if command: self.btn.after(50, command)
        item.bind("<ButtonRelease-1>", _click)
        
    def add_separator(self):
        tk.Frame(self.frame, bg="#eeeeee", height=1).pack(fill=tk.X, padx=5, pady=4)
        
    def add_cascade(self, label, values, var, suffix=""):
        item = tk.Label(self.frame, font=self.font, bg="white", anchor="w", padx=15, pady=6)
        item.pack(fill=tk.X)
        self.items.append(item)
        
        def update_parent_text(*args):
            try: item.config(text=f"{label}: {var.get()}{suffix}    ▶")
            except Exception: pass
        var.trace_add("write", update_parent_text)
        update_parent_text()
        
        submenu = tk.Toplevel(self.top)
        submenu.overrideredirect(True)
        submenu.withdraw()
        submenu.attributes('-topmost', True)
        self.submenus.append(submenu)
        
        sub_frame = tk.Frame(submenu, bg="white", highlightbackground="#cccccc", highlightthickness=1)
        sub_frame.pack(fill=tk.BOTH, expand=True)
        
        def update_sub_text(lbl, v, val): lbl.config(text=f"{'✔  ' if v.get() == val else '    '}{val}{suffix}")
            
        sub_items =[]
        for val in values:
            sub_item = tk.Label(sub_frame, font=self.font, bg="white", anchor="w", padx=15, pady=4)
            sub_item.pack(fill=tk.X)
            sub_items.append(sub_item)
            
            sub_item.bind("<ButtonRelease-1>", lambda e, v=var, vl=val:[v.set(vl), self.parent_bar.deactivate()])
            sub_item.bind("<Enter>", lambda e, w=sub_item, sibs=sub_items:[s.config(bg="white") for s in sibs] +[w.config(bg="#cce8ff")])
            var.trace_add("write", lambda *args, l=sub_item, v=var, vl=val: update_sub_text(l, v, vl))
            update_sub_text(sub_item, var, val)

        def on_item_enter(e):
            for it in self.items: it.config(bg="white")
            item.config(bg="#cce8ff")
            if self.active_submenu and self.active_submenu != submenu: self.active_submenu.withdraw()
            self.active_submenu = submenu
            self.top.update_idletasks()
            x = self.top.winfo_rootx() + self.top.winfo_width()
            y = item.winfo_rooty()
            submenu.geometry(f"+{x}+{y}")
            submenu.deiconify()
            submenu.lift()
        item.bind("<Enter>", on_item_enter)

def open_api_key_dialog():
    dialog = tk.Toplevel(root)
    dialog.title("API 키 설정 (공식 엔진 연동)")
    dialog.geometry("550x350")
    dialog.transient(root)
    dialog.grab_set()
    
    dialog.update_idletasks()
    x = root.winfo_x() + (root.winfo_width() // 2) - 225
    y = root.winfo_y() + (root.winfo_height() // 2) - 190
    dialog.geometry(f"+{x}+{y}")

    pad_frame = tk.Frame(dialog, padx=20, pady=10)
    pad_frame.pack(fill=tk.BOTH, expand=True)

    def create_api_entry(parent, label_text, default_val):
        tk.Label(parent, text=label_text, font=global_font, anchor="w").pack(fill=tk.X, pady=(10, 2))
        entry = tk.Entry(parent, width=45, font=global_font)
        entry.insert(0, default_val)
        entry.pack(fill=tk.X)
        return entry

    openai_entry = create_api_entry(pad_frame, "OpenAI API Key (gpt-4o-mini):", current_settings.get("openai_api_key", ""))
    deepl_entry = create_api_entry(pad_frame, "DeepL API Key (Free/Pro):", current_settings.get("deepl_api_key", ""))
    google_entry = create_api_entry(pad_frame, "Google Cloud Translation API Key:", current_settings.get("google_api_key", ""))
    ms_entry = create_api_entry(pad_frame, "MS Azure API Key (예: 키,지역코드 형식):", current_settings.get("ms_api_key", ""))

    def save_keys():
        current_settings["openai_api_key"] = openai_entry.get().strip()
        current_settings["deepl_api_key"] = deepl_entry.get().strip()
        current_settings["google_api_key"] = google_entry.get().strip()
        current_settings["ms_api_key"] = ms_entry.get().strip()
        save_current_config()
        messagebox.showinfo("저장 완료", "API 키가 성공적으로 적용되었습니다.", parent=dialog)
        dialog.destroy()

    tk.Button(pad_frame, text="저장 및 닫기", font=global_font, command=save_keys, bg="#4CAF50", fg="white", width=15).pack(pady=(20, 0))

def open_save_options_dialog():
    dialog = tk.Toplevel(root)
    dialog.title("문서 저장 상세 설정")
    dialog.geometry("460x440") 
    dialog.transient(root)
    dialog.grab_set()
    
    dialog.update_idletasks()
    x = root.winfo_x() + (root.winfo_width() // 2) - 230
    y = root.winfo_y() + (root.winfo_height() // 2) - 220
    dialog.geometry(f"+{x}+{y}")

    pad_frame = tk.Frame(dialog, padx=25, pady=20)
    pad_frame.pack(fill=tk.BOTH, expand=True)

    def create_spinbox_row(parent, label_text, current_val, from_val, to_val, inc_val, is_float=False):
        row = tk.Frame(parent)
        row.pack(fill=tk.X, pady=6)
        tk.Label(row, text=label_text, font=global_font, width=18, anchor="w").pack(side=tk.LEFT)
        if is_float:
            sb = ttk.Spinbox(row, from_=from_val, to=to_val, increment=inc_val, font=global_font, width=15, format="%.1f")
        else:
            sb = ttk.Spinbox(row, from_=from_val, to=to_val, increment=inc_val, font=global_font, width=15)
        sb.set(current_val)
        sb.pack(side=tk.LEFT, padx=5)
        return sb

    def create_combobox_row(parent, label_text, current_val, values):
        row = tk.Frame(parent)
        row.pack(fill=tk.X, pady=6)
        tk.Label(row, text=label_text, font=global_font, width=18, anchor="w").pack(side=tk.LEFT)
        cb = ttk.Combobox(row, values=values, state="readonly", font=global_font, width=15)
        cb.set(current_val)
        cb.pack(side=tk.LEFT, padx=5)
        return cb

    def create_checkbox_row(parent, label_text, var_obj):
        row = tk.Frame(parent)
        row.pack(fill=tk.X, pady=6)
        tk.Label(row, text=label_text, font=global_font, width=18, anchor="w").pack(side=tk.LEFT)
        chk = ttk.Checkbutton(row, variable=var_obj, text="(Word, PDF 저장 시 적용)")
        chk.pack(side=tk.LEFT, padx=5)
        return chk

    ff_entry = create_combobox_row(pad_frame, "저장 글꼴:", doc_font_family_var.get(), ALLOWED_DOC_FONTS)
    fs_entry = create_spinbox_row(pad_frame, "저장 글자 크기 (pt):", doc_font_size_var.get(), 8, 72, 1)
    ls_entry = create_spinbox_row(pad_frame, "줄 간격 (배수):", doc_line_spacing_var.get(), 1.0, 3.0, 0.1, True)
    ps_entry = create_spinbox_row(pad_frame, "문단 뒤 간격 (pt):", doc_para_spacing_var.get(), 0, 50, 1)
    mx_entry = create_spinbox_row(pad_frame, "좌우 여백 (mm):", doc_margin_x_var.get(), 5, 100, 1)
    my_entry = create_spinbox_row(pad_frame, "상하 여백 (mm):", doc_margin_y_var.get(), 5, 100, 1)
    pn_entry = create_checkbox_row(pad_frame, "하단 페이지 번호 표시:", doc_page_numbers_var)

    def save_doc_options():
        try:
            doc_font_family_var.set(ff_entry.get())
            doc_font_size_var.set(int(float(fs_entry.get())))
            doc_line_spacing_var.set(float(ls_entry.get()))
            doc_para_spacing_var.set(int(float(ps_entry.get())))
            doc_margin_x_var.set(int(float(mx_entry.get())))
            doc_margin_y_var.set(int(float(my_entry.get())))
            save_current_config()
            messagebox.showinfo("저장 완료", "문서 저장 세부 설정이 적용되었습니다.", parent=dialog)
            dialog.destroy()
        except ValueError:
            messagebox.showerror("입력 오류", "숫자 형식으로 올바르게 입력해주세요.", parent=dialog)

    tk.Button(pad_frame, text="저장 및 닫기", font=global_font, command=save_doc_options, bg="#4CAF50", fg="white", width=15).pack(pady=(25, 0))

# ----------------- GUI 메인 초기화 -----------------
root = TkinterDnD.Tk()
root.title("다중 엔진 문서 번역기 프로그램 V13 (2026.10.20)")
root.geometry("1100x800")

# --- [스타일 픽스 추가 부분] ---
style = ttk.Style()
if 'clam' in style.theme_names():
    style.theme_use('clam')  
    style.configure("TSpinbox", arrowsize=24, padding=4)  
    style.configure("TCombobox", arrowsize=20, padding=4)
    
    # 콤보박스/스핀박스 포커스 시 글자가 하얗게 변하는 현상(안 보이는 현상) 방지
    style.map("TCombobox", 
              fieldbackground=[("readonly", "white")],
              selectbackground=[("readonly", "white"), ("focus", "white")],
              selectforeground=[("readonly", "black"), ("focus", "black")],
              foreground=[("readonly", "black")])
              
    style.map("TSpinbox", 
              fieldbackground=[("readonly", "white")],
              selectbackground=[("readonly", "white"), ("focus", "white")],
              selectforeground=[("readonly", "black"), ("focus", "black")],
              foreground=[("readonly", "black")])

current_settings = load_settings()
font_size_var = tk.IntVar(value=current_settings["font_size"])
chunk_size_var = tk.IntVar(value=current_settings["chunk_size"])
engine_var = tk.StringVar(value=current_settings["engine"])
paper_size_var = tk.StringVar(value=current_settings["paper_size"])

doc_font_family_var = tk.StringVar(value=current_settings["doc_font_family"])
doc_font_size_var = tk.IntVar(value=current_settings["doc_font_size"])
doc_line_spacing_var = tk.DoubleVar(value=current_settings["doc_line_spacing"])
doc_para_spacing_var = tk.IntVar(value=current_settings["doc_para_spacing"])
doc_margin_x_var = tk.IntVar(value=current_settings["doc_margin_x"])
doc_margin_y_var = tk.IntVar(value=current_settings["doc_margin_y"])
doc_page_numbers_var = tk.BooleanVar(value=current_settings.get("doc_page_numbers", True))

global_font = tkfont.Font(family="맑은 고딕", size=font_size_var.get())
bold_font = tkfont.Font(family="맑은 고딕", size=font_size_var.get(), weight="bold")
text_font = tkfont.Font(family=doc_font_family_var.get(), size=font_size_var.get()) 

tkfont.nametofont("TkDefaultFont").configure(family="맑은 고딕", size=font_size_var.get())
tkfont.nametofont("TkMenuFont").configure(family="맑은 고딕", size=font_size_var.get())
tkfont.nametofont("TkTextFont").configure(family="맑은 고딕", size=font_size_var.get())
root.option_add("*TCombobox*Listbox.font", global_font)
root.protocol("WM_DELETE_WINDOW", save_settings)

def update_font_size(*args):
    new_size = font_size_var.get()
    global_font.configure(size=new_size)
    bold_font.configure(size=new_size)
    tkfont.nametofont("TkDefaultFont").configure(size=new_size)
    tkfont.nametofont("TkMenuFont").configure(size=new_size)
    tkfont.nametofont("TkTextFont").configure(size=new_size)
    text_font.configure(size=new_size)

def update_text_font_family(*args):
    text_font.configure(family=doc_font_family_var.get())

font_size_var.trace_add("write", update_font_size)
doc_font_family_var.trace_add("write", update_text_font_family) 

menubar = CustomMenuBar(root)

file_menu = CustomDropdownMenu(menubar, "파일(F)", global_font)
file_menu.add_command("문서 파일 불러오기...", load_from_file)
file_menu.add_command("웹사이트 텍스트로 불러오기...", load_from_url)
file_menu.add_separator()
file_menu.add_command("전체 초기화", on_reset)
file_menu.add_command("종료", save_settings)

save_menu = CustomDropdownMenu(menubar, "저장(S)", global_font)
save_menu.add_command("좌측 원문 저장", save_source_file)
save_menu.add_command("우측 번역본 저장", save_result_file)

setting_menu = CustomDropdownMenu(menubar, "설정(O)", global_font)
setting_menu.add_cascade("번역 엔진", ALLOWED_ENGINES, engine_var, "")
setting_menu.add_cascade("글자 크기 (화면용)", ALLOWED_FONT_SIZES, font_size_var, "pt")
setting_menu.add_cascade("번역 단위", ALLOWED_CHUNK_SIZES, chunk_size_var, "자")
setting_menu.add_cascade("문서 저장 기본 용지", ALLOWED_PAPER_SIZES, paper_size_var, "")
setting_menu.add_separator()
setting_menu.add_command("[ 문서 저장 상세 설정 ]", open_save_options_dialog, is_special=False)
setting_menu.add_command("[ API 키 설정 (공식 엔진) ]", open_api_key_dialog, is_special=True)

# ----------------- 리본 메뉴 -----------------
ribbon_frame = tk.Frame(root, bg="#f5f6f7", bd=1, relief=tk.RAISED, pady=6, padx=5)
ribbon_frame.pack(side=tk.TOP, fill=tk.X)

def create_ribbon_btn(parent, text, command, highlight=False):
    bg_color = "#4CAF50" if highlight else "#ffffff"
    fg_color = "white" if highlight else "black"
    btn = tk.Button(parent, text=text, font=global_font, bg=bg_color, fg=fg_color, relief=tk.RAISED if highlight else tk.GROOVE, command=command, padx=8, pady=2)
    btn.pack(side=tk.LEFT, padx=3, pady=2, ipady=3)

tk.Frame(ribbon_frame, width=2, bg="#cccccc").pack(side=tk.LEFT, fill=tk.Y, padx=10, pady=3)
create_ribbon_btn(ribbon_frame, "📁 파일 열기", load_from_file)
create_ribbon_btn(ribbon_frame, "🌐 웹 열기", load_from_url)
tk.Frame(ribbon_frame, width=2, bg="#cccccc").pack(side=tk.LEFT, fill=tk.Y, padx=10, pady=3)
create_ribbon_btn(ribbon_frame, "💾 원문 저장", save_source_file)
create_ribbon_btn(ribbon_frame, "💾 결과 저장", save_result_file)
tk.Frame(ribbon_frame, width=2, bg="#cccccc").pack(side=tk.LEFT, fill=tk.Y, padx=10, pady=3)
create_ribbon_btn(ribbon_frame, "🔄 화면 초기화", on_reset)
create_ribbon_btn(ribbon_frame, "▶ 번역 실행", on_translate, highlight=True)
create_ribbon_btn(ribbon_frame, "■ 번역 중지", stop_translation)

tk.Frame(root, height=1, bg="#d3d3d3").pack(fill=tk.X)

# ----------------- 메인 UI 레이아웃 -----------------
status_var = tk.StringVar(value="상태: 대기 중")
progress_var = tk.IntVar(value=0)

status_frame = tk.Frame(root, bd=1, relief=tk.SUNKEN)
status_frame.pack(side=tk.BOTTOM, fill=tk.X)

status_label = tk.Label(
    status_frame,
    textvariable=status_var,
    anchor=tk.W,
    font=global_font,
    padx=10,
    pady=5
)
status_label.pack(side=tk.LEFT, fill=tk.X, expand=True)

progress_text_var = tk.StringVar(value="0%")
speed_text_var = tk.StringVar(value="0.0블록/s")
remaining_time_var = tk.StringVar(value="남은 시간 --:--:--")
elapsed_time_var = tk.StringVar(value="소요 시간 00:00:00")

elapsed_time_label = tk.Label(
    status_frame,
    textvariable=elapsed_time_var,
    font=global_font,
    width=18,
    anchor=tk.E
)
elapsed_time_label.pack(side=tk.RIGHT, padx=(0, 10), pady=5)

remaining_time_label = tk.Label(
    status_frame,
    textvariable=remaining_time_var,
    font=global_font,
    width=18,
    anchor=tk.E
)
remaining_time_label.pack(side=tk.RIGHT, padx=(0, 10), pady=5)

speed_label = tk.Label(
    status_frame,
    textvariable=speed_text_var,
    font=global_font,
    width=12,
    anchor=tk.E
)
speed_label.pack(side=tk.RIGHT, padx=(0, 8), pady=5)

progress_bar = ttk.Progressbar(
    status_frame,
    variable=progress_var,
    maximum=100,
    mode="determinate",
    length=180
)
progress_bar.pack(side=tk.RIGHT, padx=(5, 8), pady=5)

progress_label = tk.Label(
    status_frame,
    textvariable=progress_text_var,
    font=global_font,
    width=6,
    anchor=tk.E
)
progress_label.pack(side=tk.RIGHT, padx=(0, 5), pady=5)


def format_seconds(seconds):
    try:
        seconds = int(seconds)
        h = seconds // 3600
        m = (seconds % 3600) // 60
        s = seconds % 60
        return f"{h:02d}:{m:02d}:{s:02d}"
    except Exception:
        return "--:--:--"


def update_progress_text(*args):
    try:
        maximum = progress_bar.cget("maximum")
        current = progress_var.get()

        if maximum <= 0:
            percent = 0
        else:
            percent = int((current / maximum) * 100)

        progress_text_var.set(f"{percent}%")
    except Exception:
        progress_text_var.set("0%")

progress_var.trace_add("write", update_progress_text)

top_frame = tk.Frame(root, pady=10, padx=10)
top_frame.pack(fill=tk.X)

tk.Label(top_frame, text="번역 엔진:", font=global_font).pack(side=tk.LEFT, padx=(0, 5))
engine_combobox = ttk.Combobox(top_frame, values=ALLOWED_ENGINES, state="readonly", width=16, font=global_font)
engine_combobox.set(engine_var.get())
engine_combobox.pack(side=tk.LEFT, padx=(0, 15))
engine_combobox.bind("<<ComboboxSelected>>", lambda e: engine_var.set(engine_combobox.get()))
engine_var.trace_add("write", lambda *a: engine_combobox.set(engine_var.get()))

tk.Label(top_frame, text="도착 언어:", font=global_font).pack(side=tk.LEFT, padx=(0, 5))
lang_combobox = ttk.Combobox(top_frame, values=list(LANGUAGES.keys()), state="readonly", width=12, font=global_font)
lang_combobox.set(current_settings["target_lang"]) 
lang_combobox.pack(side=tk.LEFT, padx=(0, 15))

tk.Label(top_frame, text="단위:", font=global_font).pack(side=tk.LEFT, padx=(0, 5))
chunk_combobox = ttk.Combobox(top_frame, values=[f"{c}자" for c in ALLOWED_CHUNK_SIZES], state="readonly", width=6, font=global_font)
chunk_combobox.set(f"{current_settings['chunk_size']}자")
chunk_combobox.pack(side=tk.LEFT, padx=(0, 15))
chunk_combobox.bind("<<ComboboxSelected>>", lambda e: chunk_size_var.set(int(chunk_combobox.get().replace("자", ""))))
chunk_size_var.trace_add("write", lambda *a: chunk_combobox.set(f"{chunk_size_var.get()}자"))

tk.Label(top_frame, text="출판 용지:", font=global_font).pack(side=tk.LEFT, padx=(0, 5))
paper_combobox = ttk.Combobox(top_frame, values=ALLOWED_PAPER_SIZES, state="readonly", width=18, font=global_font)
paper_combobox.set(paper_size_var.get())
paper_combobox.pack(side=tk.LEFT, padx=(0, 5))
paper_combobox.bind("<<ComboboxSelected>>", lambda e: paper_size_var.set(paper_combobox.get()))
paper_size_var.trace_add("write", lambda *a: paper_combobox.set(paper_size_var.get()))

# ----------------- 텍스트 입력 창 및 동기화 -----------------
is_syncing = False
def sync_scroll_input(*args):
    global is_syncing
    input_scroll.set(*args)
    if not is_syncing: 
        is_syncing = True
        text_output.yview_moveto(args[0])
        is_syncing = False

def sync_scroll_output(*args):
    global is_syncing
    output_scroll.set(*args)
    if not is_syncing: 
        is_syncing = True
        text_input.yview_moveto(args[0])
        is_syncing = False

def scroll_both_bars(*args):
    text_input.yview(*args)
    text_output.yview(*args)

main_frame = tk.Frame(root)
main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 10))
main_frame.columnconfigure(0, weight=1, uniform="equal")
main_frame.columnconfigure(1, weight=1, uniform="equal")
main_frame.rowconfigure(0, weight=1)

novel_view_config = {"wrap": tk.WORD, "spacing1": 8, "spacing2": 5, "spacing3": 12, "padx": 25, "pady": 20, "bg": "#FDFCF7"}

input_frame = tk.Frame(main_frame)
input_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 5))
tk.Label(input_frame, text="[ 원문 (다중 파일 드래그 / 편집 가능) ]", anchor="w", font=global_font).pack(fill=tk.X, pady=(0, 5))
input_scroll = tk.Scrollbar(input_frame)
input_scroll.pack(side=tk.RIGHT, fill=tk.Y)
text_input = tk.Text(input_frame, font=text_font, **novel_view_config)
text_input.pack(fill=tk.BOTH, expand=True)
getattr(text_input, "drop_target_register")(DND_FILES)
getattr(text_input, "dnd_bind")("<<Drop>>", handle_drop)

output_frame = tk.Frame(main_frame)
output_frame.grid(row=0, column=1, sticky="nsew", padx=(5, 0))
tk.Label(output_frame, text="[ 번역 결과 ]", anchor="w", font=global_font).pack(fill=tk.X, pady=(0, 5))
output_scroll = tk.Scrollbar(output_frame)
output_scroll.pack(side=tk.RIGHT, fill=tk.Y)
text_output = tk.Text(output_frame, font=text_font, **novel_view_config) 
text_output.pack(fill=tk.BOTH, expand=True)

text_input.config(yscrollcommand=sync_scroll_input)
text_output.config(yscrollcommand=sync_scroll_output)
input_scroll.config(command=scroll_both_bars)
output_scroll.config(command=scroll_both_bars)

# 하이라이트 동기화
sync_timer = None

def sync_highlight(source, target):
    try:
        start_index = source.index(tk.SEL_FIRST)
        end_index = source.index(tk.SEL_LAST)
        start_line, start_col = map(int, start_index.split('.'))
        end_line, end_col = map(int, end_index.split('.'))
        
        target.tag_remove("sync_hl", "1.0", tk.END)
        limit_end = min(end_line, start_line + 200)
        
        for line in range(start_line, limit_end + 1):
            src_len = len(source.get(f"{line}.0", f"{line}.end"))
            tgt_len = len(target.get(f"{line}.0", f"{line}.end"))
            if src_len == 0 or tgt_len == 0: continue
            
            c1 = start_col if line == start_line else 0
            c2 = end_col if line == end_line else src_len
            if line == limit_end and limit_end < end_line: c2 = src_len 
                
            out_c1 = int((c1 / src_len) * tgt_len)
            out_c2 = int((c2 / src_len) * tgt_len)
            if out_c1 == out_c2 and c1 != c2: out_c2 += 1
                
            target.tag_add("sync_hl", f"{line}.{out_c1}", f"{line}.{out_c2}")
            
        target.tag_config("sync_hl", background="#DDEBF7")
    except tk.TclError:
        target.tag_remove("sync_hl", "1.0", tk.END)

def on_sync_out(e):
    global sync_timer
    if sync_timer: root.after_cancel(sync_timer)
    sync_timer = root.after(100, lambda: sync_highlight(text_input, text_output))

def on_sync_in(e):
    global sync_timer
    if sync_timer: root.after_cancel(sync_timer)
    sync_timer = root.after(100, lambda: sync_highlight(text_output, text_input))

text_input.bind("<<Selection>>", on_sync_out)
text_output.bind("<<Selection>>", on_sync_in)

root.mainloop()
