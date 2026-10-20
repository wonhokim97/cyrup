# -*- coding: utf-8 -*-
"""PDF Master Pro V9 - 타입 명시 및 작업 안정성 보완본.

필수: PySide6 pymupdf python-docx python-pptx openpyxl pillow
선택: ebooklib (EPUB), comtypes (Windows MS Office 불러오기)
실행: python PDF_Master_Pro_V9_fixed.py

Windows / MS Office COM 동작은 해당 환경에서 별도 확인이 필요합니다.
"""
from __future__ import annotations

import sys
import os
import io
import re
import html
import uuid
import datetime
import tempfile
from collections import OrderedDict
from dataclasses import dataclass, field
from typing import Any, Optional

import pymupdf as fitz
from PySide6.QtCore import (
    Qt, QObject, Signal, Slot, QRunnable, QThreadPool, QTimer,
    QSettings, QSize, QByteArray, QLoggingCategory,
)
from PySide6.QtGui import (
    QAction, QActionGroup, QPixmap, QImage, QIcon, QFont,
    QFontDatabase, QFontMetrics, QCursor,
)
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QFileDialog, QVBoxLayout,
    QHBoxLayout, QPushButton, QListWidget, QListWidgetItem, QLabel,
    QScrollArea, QMenu, QAbstractItemView, QSplitter, QFrame,
    QMessageBox, QProgressBar, QLineEdit, QSizePolicy, QDialog, QTextEdit,
)

ROLE = Qt.ItemDataRole.UserRole
PageKey = tuple[int, int, int]


def resource_base_dir() -> str:
    base = getattr(sys, "_MEIPASS", None)
    return str(base) if base else os.path.dirname(os.path.abspath(__file__))


def read_text(path: str) -> str:
    raw = open(path, "rb").read()
    for encoding in ("utf-8-sig", "utf-16", "cp949", "euc-kr"):
        if encoding == "utf-16" and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        try:
            return raw.decode(encoding)
        except UnicodeError:
            continue
    return raw.decode("latin1")


@dataclass
class SourceDoc:
    doc: fitz.Document
    name: str
    path: str
    source_type: str
    render_path: str
    temp_paths: list[str] = field(default_factory=list)

    def close(self) -> None:
        self.doc.close()

    def cleanup(self) -> None:
        for path in self.temp_paths:
            try:
                os.remove(path)
            except OSError:
                pass
        self.temp_paths.clear()


class ResponsiveLabel(QLabel):
    def __init__(self, parent: Optional[QWidget] = None):
        super().__init__(parent)
        self._current_pixmap: Optional[QPixmap] = None
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setMinimumSize(0, 0)
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.clear()

    def clear(self) -> None:
        self._current_pixmap = None
        super().clear()
        self.setText("파일을 로드하고 리스트에 마우스를 올리세요")

    def set_pixmap_fixed(self, pixmap: QPixmap) -> None:
        self._current_pixmap = pixmap
        self.update_display()

    def update_display(self) -> None:
        if self._current_pixmap is not None:
            super().setPixmap(self._current_pixmap.scaled(
                self.size(), Qt.AspectRatioMode.KeepAspectRatio,
                Qt.TransformationMode.SmoothTransformation,
            ))

    def resizeEvent(self, event):
        self.update_display()
        super().resizeEvent(event)


class RenderSignals(QObject):
    ready = Signal(object)
    error = Signal(object)


def build_page(out: fitz.Document, source: fitz.Document, index: int,
               rotation: int, size: tuple[float, float],
               align: str, margins: bool, raster: bool = False,
               dpi: int = 150, quality: int = 85) -> None:
    """프리뷰/병합이 공유하는 배치 엔진. 단일 페이지의 회전을 먼저 정규화."""
    with fitz.open() as normalized:
        normalized.insert_pdf(source, from_page=index, to_page=index)
        page = normalized[0]
        page.set_rotation((page.rotation + rotation) % 360)
        width, height = size
        target = out.new_page(width=width, height=height)
        rect = target.rect
        if margins:
            margin = min(25.0, width / 10, height / 10)
            area = fitz.Rect(margin, margin, width - margin, height - margin)
            scale = min(area.width / page.rect.width, area.height / page.rect.height)
            dw, dh = page.rect.width * scale, page.rect.height * scale
            x = area.x0 + (area.width - dw) / 2
            y = area.y0 if align == "top" else area.y0 + (area.height - dh) / 2
            rect = fitz.Rect(x, y, x + dw, y + dh)
        if raster:
            pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csRGB, alpha=False)
            target.insert_image(rect, stream=pix.tobytes("jpeg", jpg_quality=quality))
        else:
            target.show_pdf_page(rect, normalized, 0, keep_proportion=True)


class RenderTask(QRunnable):
    def __init__(self, signals: RenderSignals, payload: dict[str, Any]):
        super().__init__()
        self.signals = signals
        self.payload = payload

    def run(self) -> None:
        p = self.payload
        try:
            with fitz.open(p["path"]) as source:
                if p["kind"] == "preview":
                    with fitz.open() as preview:
                        build_page(preview, source, p["index"], p["rotation"],
                                   p["size"], p["align"], p["margins"])
                        page = preview[0]
                        pix = page.get_pixmap(
                            matrix=fitz.Matrix(p["width"] / page.rect.width,
                                               p["width"] / page.rect.width),
                            colorspace=fitz.csRGB, alpha=False)
                else:
                    page = source[p["index"]]
                    scale = p["width"] / page.rect.width
                    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale),
                                          colorspace=fitz.csRGB, alpha=False)
                self.signals.ready.emit({
                    **p, "samples": bytes(pix.samples), "w": pix.width,
                    "h": pix.height, "stride": pix.stride,
                })
        except Exception as exc:
            self.signals.error.emit({**p, "message": str(exc)})


class MultiSelectListWidget(QListWidget):
    def __init__(self, parent: IntegratedPDFMaster, is_work: bool = False):
        super().__init__(parent)
        self.main_win: IntegratedPDFMaster = parent
        self.is_work = is_work
        self.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.setMouseTracking(True)
        self.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.customContextMenuRequested.connect(self.show_context_menu)
        self.setMinimumWidth(0)
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.setAcceptDrops(True)
        self.viewport().setAcceptDrops(True)
        self.setDropIndicatorShown(True)
        self.setDragEnabled(is_work)
        self.setDragDropMode(
            QAbstractItemView.DragDropMode.InternalMove if is_work
            else QAbstractItemView.DragDropMode.DropOnly)
        self.setDefaultDropAction(
            Qt.DropAction.MoveAction if is_work else Qt.DropAction.CopyAction)
        self.setStyleSheet(
            "QListWidget::item {padding:0px;margin:0px;}"
            "QListWidget::item:selected {background:#d0ebff;}")

    def mouseMoveEvent(self, event):
        item = self.itemAt(event.position().toPoint())
        if item is not None:
            self.main_win.handle_item_preview(item)
        super().mouseMoveEvent(event)

    def wheelEvent(self, event):
        if not self.count():
            super().wheelEvent(event)
            return
        row = self.currentRow()
        row += -1 if event.angleDelta().y() > 0 else 1
        row = max(0, min(self.count() - 1, row))
        self.setCurrentRow(row)
        self.main_win.handle_item_preview(self.item(row))
        event.accept()

    def dragEnterEvent(self, event):
        if self.is_work:
            if event.source() is self:
                super().dragEnterEvent(event)
            else:
                event.ignore()
        elif event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        if self.is_work:
            if event.source() is self:
                super().dragMoveEvent(event)
            else:
                event.ignore()
        elif event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        if self.is_work:
            if event.source() is self:
                super().dropEvent(event)
                self.main_win.on_worklist_reordered()
            else:
                event.ignore()
        elif event.mimeData().hasUrls():
            paths = [url.toLocalFile() for url in event.mimeData().urls()]
            self.main_win.load_files(paths, source="드롭")
            event.acceptProposedAction()
        else:
            event.ignore()

    def show_context_menu(self, pos) -> None:
        item = self.itemAt(pos)
        if item is None:
            return
        if not item.isSelected():
            self.setCurrentItem(item)
        menu = QMenu(self)
        if not self.is_work:
            menu.addAction("선택 항목 모두 추가").triggered.connect(
                self.main_win.add_selected_to_work)
        else:
            menu.addAction("반시계 90도 회전").triggered.connect(
                lambda: self.main_win.rotate_work_items(270))
            menu.addAction("시계 90도 회전").triggered.connect(
                lambda: self.main_win.rotate_work_items(90))
            menu.addAction("회전 초기화").triggered.connect(
                lambda: self.main_win.rotate_work_items(0, reset=True))
            menu.addSeparator()
            menu.addAction("선택 항목 삭제").triggered.connect(
                self.main_win.remove_work_items)
        menu.exec(QCursor.pos())


class TextImportDialog(QDialog):
    def __init__(self, parent: IntegratedPDFMaster, txt_path: str):
        super().__init__(parent)
        self.main_win: IntegratedPDFMaster = parent
        self.txt_path = txt_path
        self.saved_pdf: Optional[str] = None
        self.setWindowTitle(f"텍스트 불러오기 - {os.path.basename(txt_path)}")
        self.resize(800, 650)
        layout = QVBoxLayout(self)
        layout.addWidget(QLabel("원문 확인 및 수정 — 설정한 소설책 규격으로 저장"))
        self.text_edit = QTextEdit()
        self.text_edit.setPlainText(read_text(txt_path))
        layout.addWidget(self.text_edit)
        buttons = QHBoxLayout()
        for title, callback in (
            ("PDF 저장 후 리스트 추가", self.save_pdf),
            ("EPUB 저장 (단독)", self.save_epub),
            ("취소", self.reject),
        ):
            button = QPushButton(title)
            button.clicked.connect(callback)
            buttons.addWidget(button)
        layout.addLayout(buttons)

    def save_pdf(self) -> None:
        path = self.main_win.pick_save("소설변환", "pdf", parent=self)
        if not path:
            return
        try:
            with self.main_win.create_novel_pdf_from_text(
                    self.text_edit.toPlainText()) as doc:
                self.main_win._save_pdf(doc, path)
            self.saved_pdf = path
            self.accept()
        except Exception as exc:
            QMessageBox.warning(self, "PDF 저장 오류", str(exc))

    def save_epub(self) -> None:
        try:
            from ebooklib import epub
            path = self.main_win.pick_save("소설변환", "epub", parent=self)
            if not path:
                return
            book = epub.EpubBook()
            book.set_identifier(uuid.uuid4().hex)
            book.set_title(os.path.splitext(os.path.basename(self.txt_path))[0])
            book.set_language("ko")
            chapter = epub.EpubHtml(title="본문", file_name="content.xhtml", lang="ko")
            chapter.content = "".join(
                f"<p>{html.escape(line.strip())}</p>" if line.strip()
                else "<p><br/></p>"
                for line in self.text_edit.toPlainText().splitlines())
            style = epub.EpubItem(
                uid="style", file_name="style/book.css", media_type="text/css",
                content=b"body{margin:5%;font-family:sans-serif}"
                        b"p{text-indent:1em;line-height:1.8;margin:.3em 0}")
            book.add_item(style)
            chapter.add_item(style)
            book.add_item(chapter)
            book.toc = [chapter]
            book.spine = ["nav", chapter]
            book.add_item(epub.EpubNcx())
            book.add_item(epub.EpubNav())
            epub.write_epub(path, book, {})
            self.main_win.set_status("EPUB 저장 완료")
            self.accept()
        except ImportError:
            QMessageBox.warning(self, "패키지 필요", "pip install ebooklib")
        except Exception as exc:
            QMessageBox.warning(self, "EPUB 저장 오류", str(exc))


class IntegratedPDFMaster(QMainWindow):
    SUPPORTED_EXTS = {"pdf", "txt", "jpg", "jpeg", "png", "bmp", "docx", "pptx", "xlsx"}
    PREVIEW_CACHE_LIMIT = 80
    THUMB_CACHE_LIMIT = 2500
    THUMB_QUEUE_LIMIT_PER_LOAD = 450
    PREVIEW_MAX_W = 1200
    PREVIEW_DEBOUNCE_MS = 80
    TEXT_MODE_PLAIN = "plain"
    TEXT_MODE_LAYOUT = "layout"
    CONVERT_MODE_AUTO = "auto"
    CONVERT_MODE_IMAGE = "image"
    CONVERT_MODE_MIXED = "mixed"
    MERGE_SIZE_KEEP = "keep"
    MERGE_SIZE_FIRST = "first"
    MERGE_SIZE_A4 = "a4"
    MERGE_ALIGN_CENTER = "center"
    MERGE_ALIGN_TOP = "top"
    MERGE_ROTATE_KEEP = "keep"
    MERGE_ROTATE_PORTRAIT = "portrait"
    MERGE_ROTATE_LANDSCAPE = "landscape"
    PDF_SAVE_DEFAULT = "default"
    PDF_SAVE_SMALL = "small"
    PDF_SAVE_PRINT = "print"

    def __init__(self):
        super().__init__()
        self.setWindowTitle("통합 PDF 마스터 Pro V9 — 수정본")
        self.resize(1400, 850)
        self.settings = QSettings("MySoft", "PDFMaster_AllInOne_Final")
        self.docs: list[SourceDoc] = []
        self.last_path = self._setting_str("last_path", os.path.expanduser("~"))
        self.current_dpi = self._setting_int("dpi", 150, {72, 150, 300, 600})
        self.current_font_size = self._setting_int("font_size", 11, {10, 11, 12, 14, 16})
        self.text_import_font_size = self._setting_int(
            "text_import_font_size", 11, {9, 10, 11, 12, 13, 14})
        for key, default in (
            ("text_mode", "plain"), ("convert_mode", "auto"),
            ("novel_size_mode", "A5"), ("merge_size_mode", "keep"),
            ("merge_align_mode", "center"), ("merge_rotate_mode", "keep"),
            ("merge_orient_mode", "keep"), ("pdf_save_mode", "default"),
        ):
            setattr(self, key, self._setting_str(key, default))
        self.text_mode: str
        self.convert_mode: str
        self.novel_size_mode: str
        self.merge_size_mode: str
        self.merge_align_mode: str
        self.merge_rotate_mode: str
        self.merge_orient_mode: str
        self.pdf_save_mode: str
        self.pdf_rasterize_mode = (
            self._setting_str("pdf_rasterize_mode", "false").lower()
            in {"true", "1", "yes"})
        self.font_path = self._find_korean_font_path()
        self._preview_cache: OrderedDict[tuple[Any, ...], QPixmap] = OrderedDict()
        self._thumb_cache: OrderedDict[tuple[int, int], QPixmap] = OrderedDict()
        self._orig_item_map: dict[tuple[int, int], QListWidgetItem] = {}
        self._inflight: set[tuple[Any, ...]] = set()
        self._work_key_set: set[tuple[int, int]] = set()
        self._generation = 0
        self._thumb_generation = 0
        self._icon_px = 48
        self._pending_preview_key: Optional[PageKey] = None
        self._current_preview_key: Optional[tuple[Any, ...]] = None
        self._busy = False
        self._closing = False
        self.thread_pool = QThreadPool(self)
        # PyMuPDF 동시 스레드 호출을 피하도록 렌더 작업은 직렬화합니다.
        # 문서 변경/변환 작업 전에도 렌더 풀 완료를 기다립니다.
        self.thread_pool.setMaxThreadCount(1)
        self.render_signals = RenderSignals(self)
        self.render_signals.ready.connect(self._on_render_ready)
        self.render_signals.error.connect(self._on_render_error)
        self._preview_timer = QTimer(self)
        self._preview_timer.setSingleShot(True)
        self._preview_timer.timeout.connect(self._render_pending_preview)
        self.init_ui()
        self.create_full_menu_bar()
        self.load_settings()

    def _setting_str(self, key: str, default: str) -> str:
        value = self.settings.value(key, default)
        return value if isinstance(value, str) else str(value)

    def _setting_int(self, key: str, default: int,
                     allowed: Optional[set[int]] = None) -> int:
        try:
            value = int(str(self.settings.value(key, default)))
            return default if allowed is not None and value not in allowed else value
        except (ValueError, TypeError):
            return default

    def _ts(self) -> str:
        return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

    def default_save_path(self, base_name: str, ext: str) -> str:
        return os.path.join(self.last_path, f"{base_name}_{self._ts()}.{ext}")

    def update_last_path(self, path: str) -> None:
        self.last_path = path if os.path.isdir(path) else os.path.dirname(path)
        self.settings.setValue("last_path", self.last_path)

    def pick_save(self, name: str, ext: str,
                  parent: Optional[QWidget] = None) -> str:
        path, _ = QFileDialog.getSaveFileName(
            parent or self, f"{ext.upper()} 저장",
            self.default_save_path(name, ext), f"{ext.upper()} (*.{ext})")
        if path:
            if not path.lower().endswith("." + ext):
                path += "." + ext
            self.update_last_path(path)
        return path

    def set_status(self, text: str) -> None:
        self.status_bar.setText(" ▶ " + text)

    def _find_korean_font_path(self) -> Optional[str]:
        base = resource_base_dir()
        candidates = [
            os.path.join(base, "fonts", "NotoSansKR-Regular.ttf"),
            os.path.join(base, "fonts", "NotoSansKR-Regular.otf"),
            os.path.join(base, "fonts", "NanumGothic.ttf"),
            r"C:\Windows\Fonts\malgun.ttf",
            "/System/Library/Fonts/Supplemental/AppleGothic.ttf",
            "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        ]
        return next((path for path in candidates if os.path.isfile(path)), None)

    def init_ui(self) -> None:
        """원본 화면 구성 복원: 고정 도구/상태 영역과 확장되는 본문."""
        self.central_widget = QWidget(self)
        self.setCentralWidget(self.central_widget)
        self.main_layout = QVBoxLayout(self.central_widget)
        self.main_layout.setContentsMargins(0, 0, 0, 0)
        self.main_layout.setSpacing(0)

        self.tool_panel = QFrame()
        self.tool_panel.setFixedHeight(70)
        self.tool_panel.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.tool_panel.setStyleSheet(
            "QFrame {background:#f8f9fa;border-bottom:1px solid #dee2e6;}")
        self.tool_layout = QHBoxLayout(self.tool_panel)
        self.tool_layout.setContentsMargins(9, 6, 9, 6)
        self.tool_layout.setSpacing(6)
        entries = [
            ("열기", "📂", lambda: self.open_files(True)),
            ("추가", "➕", lambda: self.open_files(False)),
            ("초기화", "🧹", self.clear_all),
            ("|", "", None),
            ("병합", "🔗", self.action_merge),
            ("추출", "✂", self.action_extract),
            ("|", "", None),
            ("Word", "📝", self.convert_to_word),
            ("PPT", "📊", self.convert_to_ppt),
            ("|", "", None),
            ("Excel", "📑", self.convert_to_excel),
            ("Text", "📄", self.convert_to_text),
            ("EPUB", "📖", self.convert_to_epub),
            ("이미지", "🖼", self.convert_to_images),
        ]
        for title, symbol, callback in entries:
            if title == "|":
                line = QFrame()
                line.setFrameShape(QFrame.Shape.VLine)
                line.setFixedSize(1, 32)
                line.setStyleSheet("background:#ced4da;border:0;")
                self.tool_layout.addWidget(line)
                continue
            button = QPushButton(symbol + "\n" + title)
            button.setFixedSize(68, 48)
            button.setStyleSheet(
                "QPushButton {background:transparent;border:0;"
                "border-radius:5px;font-weight:bold;}"
                "QPushButton:hover {background:#e9ecef;}"
                "QPushButton:pressed {background:#dee2e6;}")
            if callback is not None:
                button.clicked.connect(callback)
            self.tool_layout.addWidget(button)
        self.tool_layout.addStretch(1)
        self.main_layout.addWidget(self.tool_panel, 0)

        self.content_splitter = QSplitter(Qt.Orientation.Horizontal)
        self.content_splitter.setHandleWidth(6)
        self.content_splitter.setChildrenCollapsible(True)
        self.content_splitter.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)

        for title, is_work in (
                ("📂 원본 페이지", False), ("📑 작업 리스트", True)):
            if not is_work:
                title = "📂 원본 페이지"
            panel = QWidget()
            panel.setMinimumWidth(0)
            panel.setSizePolicy(
                QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
            pl = QVBoxLayout(panel)
            pl.setContentsMargins(9, 9, 9, 9)
            pl.setSpacing(6)
            pl.addWidget(QLabel("<b>" + title + "</b>"), 0)
            search = QLineEdit()
            search.setPlaceholderText("검색(필터): 파일명 / 페이지 ...")
            search.setClearButtonEnabled(True)
            pl.addWidget(search, 0)
            lw = MultiSelectListWidget(self, is_work)
            lw.itemClicked.connect(self.handle_item_preview)
            search.textChanged.connect(
                lambda text, target=lw: self._apply_filter(target, text))
            pl.addWidget(lw, 1)
            if is_work:
                self.work_panel = panel
                self.work_list = lw
                self.work_search = search
                lw.itemDoubleClicked.connect(self.remove_work_items)
                button = QPushButton("리스트에서 삭제")
                button.setStyleSheet(
                    "QPushButton {background:#ff4757;color:white;"
                    "font-weight:bold;border:1px solid #ef3340;"
                    "border-radius:2px;padding:3px;}"
                    "QPushButton:hover {background:#ff6b81;}")
                button.clicked.connect(self.remove_work_items)
                pl.addWidget(button, 0)
                pl.addWidget(QLabel(
                    "<span style='color:#666'>"
                    "TIP: 마우스 드래그로 순서를 바꿀 수 있습니다.</span>"), 0)
            else:
                self.orig_panel = panel
                self.orig_list = lw
                self.orig_search = search
                lw.itemDoubleClicked.connect(self.add_selected_to_work)
            self.content_splitter.addWidget(panel)

        preview = QWidget()
        preview.setMinimumWidth(0)
        preview.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        pl = QVBoxLayout(preview)
        pl.setContentsMargins(9, 9, 9, 9)
        pl.setSpacing(6)
        pl.addWidget(QLabel("<b>🔍 실시간 미리보기</b>"), 0)
        self.preview_label = ResponsiveLabel()
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setMinimumSize(0, 0)
        scroll.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        scroll.setWidget(self.preview_label)
        pl.addWidget(scroll, 1)
        self.content_splitter.addWidget(preview)
        for index, factor in enumerate((1, 1, 2)):
            self.content_splitter.setStretchFactor(index, factor)
        self.main_layout.addWidget(self.content_splitter, 1)

        self.status_container = QFrame()
        self.status_container.setFixedHeight(30)
        self.status_container.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.status_container.setStyleSheet("background:#343a40;")
        self.status_layout = QHBoxLayout(self.status_container)
        self.status_layout.setContentsMargins(10, 0, 10, 0)
        self.status_layout.setSpacing(8)
        self.status_bar = QLabel(" ▶ 시스템 준비 완료")
        self.status_bar.setStyleSheet("color:#00ff88;")
        self.status_layout.addWidget(self.status_bar, 1)
        self.progress_bar = QProgressBar()
        self.progress_bar.setFixedSize(260, 15)
        self.progress_bar.setTextVisible(True)
        self.progress_bar.setStyleSheet(
            "QProgressBar {background:#495057;border:0;border-radius:5px;"
            "color:white;text-align:center;}"
            "QProgressBar::chunk {background:#00ff88;border-radius:5px;}")
        self.progress_bar.hide()
        self.status_layout.addWidget(self.progress_bar, 0)
        self.main_layout.addWidget(self.status_container, 0)

    def _apply_filter(self, lw: QListWidget, text: str) -> None:
        for index in range(lw.count()):
            item = lw.item(index)
            if item is not None:
                item.setHidden(text.strip().lower() not in item.text().lower())

    def _option_menu(self, parent: QMenu, title: str, attribute: str,
                     choices: list[tuple[str, Any]]) -> None:
        menu = parent.addMenu(title)
        group = QActionGroup(menu)
        group.setExclusive(True)
        for label, value in choices:
            action = QAction(label, menu, checkable=True)
            action.setChecked(getattr(self, attribute) == value)
            action.triggered.connect(
                lambda checked=False, attr=attribute, val=value:
                self._change_option(attr, val))
            group.addAction(action)
            menu.addAction(action)

    def create_full_menu_bar(self) -> None:
        menu = self.menuBar().addMenu("파일")
        menu.addAction("새로 열기").triggered.connect(lambda: self.open_files(True))
        menu.addAction("파일 추가").triggered.connect(lambda: self.open_files(False))
        menu.addAction("전체 초기화").triggered.connect(self.clear_all)
        menu.addAction("종료").triggered.connect(self.close)
        pdf = self.menuBar().addMenu("PDF 생성")
        pdf.addAction("작업 리스트 병합").triggered.connect(self.action_merge)
        pdf.addAction("작업 리스트 개별 추출").triggered.connect(self.action_extract)
        self.action_rasterize_pdf = QAction(
            "DPI 강제 적용 (이미지화 / 텍스트 검색 불가)", self, checkable=True)
        self.action_rasterize_pdf.setChecked(self.pdf_rasterize_mode)
        self.action_rasterize_pdf.triggered.connect(
            lambda checked: self._change_option("pdf_rasterize_mode", checked))
        pdf.addAction(self.action_rasterize_pdf)
        convert = self.menuBar().addMenu("변환")
        for label, callback in (
            ("Word", self.convert_to_word), ("PPT", self.convert_to_ppt),
            ("Excel", self.convert_to_excel), ("Text", self.convert_to_text),
            ("EPUB", self.convert_to_epub), ("이미지", self.convert_to_images),
        ):
            convert.addAction(label).triggered.connect(callback)
        settings = self.menuBar().addMenu("설정")
        self._option_menu(settings, "Text 소설책 규격", "novel_size_mode",
                          [(x, x) for x in ("A5", "신국판", "46판", "B5")])
        self._option_menu(settings, "Text 가져오기 글자 크기", "text_import_font_size",
                          [(f"{x}pt", x) for x in (9, 10, 11, 12, 13, 14)])
        self._option_menu(settings, "출력 해상도", "current_dpi",
                          [(f"{x} DPI", x) for x in (72, 150, 300, 600)])
        self._option_menu(settings, "UI 글자 크기", "current_font_size",
                          [(f"{x}pt", x) for x in (10, 11, 12, 14, 16)])
        self._option_menu(settings, "텍스트 추출", "text_mode",
                          [("일반", "plain"), ("읽기 순서 정렬", "layout")])
        self._option_menu(settings, "변환 모드", "convert_mode",
                          [("자동 (텍스트 우선)", "auto"), ("이미지", "image"),
                           ("혼합", "mixed")])
        self._option_menu(settings, "PDF 저장", "pdf_save_mode",
                          [("일반", "default"), ("저용량", "small"),
                           ("인쇄용", "print")])
        merge = settings.addMenu("PDF 병합 옵션")
        self._option_menu(merge, "페이지 크기", "merge_size_mode",
                          [("원본 유지", "keep"), ("첫 페이지 기준", "first"),
                           ("A4", "a4")])
        self._option_menu(merge, "용지 방향", "merge_orient_mode",
                          [("원본 유지", "keep"), ("세로", "portrait"),
                           ("가로", "landscape")])
        self._option_menu(merge, "페이지 배치", "merge_align_mode",
                          [("중앙", "center"), ("상단", "top")])
        self._option_menu(merge, "내용 자동 회전", "merge_rotate_mode",
                          [("원본 유지", "keep"), ("세로로 통일", "portrait"),
                           ("가로로 통일", "landscape")])

    def _change_option(self, attribute: str, value: Any) -> None:
        setattr(self, attribute, value)
        if attribute == "current_font_size":
            self.change_ui_font(int(value))
        if attribute.startswith("merge_"):
            self.refresh_current_preview()
        self.set_status(f"설정 변경: {value}")

    def change_ui_font(self, size: int) -> None:
        self.current_font_size = size
        font = QFont("Malgun Gothic", size) if (
            "Malgun Gothic" in QFontDatabase.families()) else QApplication.font()
        font.setPointSize(size)
        QApplication.setFont(font)
        self._update_list_icon_size()

    def _update_list_icon_size(self) -> None:
        size = max(24, int(QFontMetrics(QApplication.font()).height() * 1.55))
        if size != self._icon_px:
            self._thumb_generation += 1
            self._thumb_cache.clear()
        self._icon_px = size
        for lw in (self.orig_list, self.work_list):
            lw.setFont(QApplication.font())
            lw.setIconSize(QSize(size, size))
            for index in range(lw.count()):
                item = lw.item(index)
                if item is not None:
                    item.setSizeHint(QSize(0, size + 2))

    def _begin(self, total: int, text: str) -> bool:
        if self._busy:
            return False
        self._busy = True
        self._preview_timer.stop()
        self.thread_pool.clear()
        self.thread_pool.waitForDone()
        self._inflight.clear()
        self.tool_panel.setEnabled(False)
        self.menuBar().setEnabled(False)
        self.content_splitter.setEnabled(False)
        self.progress_bar.setRange(0, max(1, total))
        self.progress_bar.setValue(0)
        self.progress_bar.setFormat(text + " %v/%m")
        self.progress_bar.show()
        return True

    def _advance(self, value: int) -> None:
        self.progress_bar.setValue(value)
        QApplication.processEvents()

    def _end(self) -> None:
        self._busy = False
        self.tool_panel.setEnabled(True)
        self.menuBar().setEnabled(True)
        self.content_splitter.setEnabled(True)
        self.progress_bar.hide()

    def _get_jpg_quality(self) -> int:
        return {"small": 70, "print": 95}.get(self.pdf_save_mode, 85)

    def _save_pdf(self, doc: fitz.Document, save_path: str) -> None:
        # 같은 경로의 원본 PDF가 열려 있어도 직접 덮어쓰지 않습니다.
        target = os.path.normcase(os.path.abspath(save_path))
        if any(os.path.normcase(os.path.abspath(sd.render_path)) == target
               for sd in self.docs):
            raise ValueError("열려 있는 원본 파일과 다른 이름으로 저장하세요.")
        folder = os.path.dirname(os.path.abspath(save_path))
        fd, temp = tempfile.mkstemp(suffix=".pdf", dir=folder)
        os.close(fd)
        try:
            doc.save(temp, garbage=4, deflate=self.pdf_save_mode != "print")
            os.replace(temp, save_path)
        finally:
            if os.path.exists(temp):
                os.remove(temp)

    def open_files(self, clear: bool) -> None:
        if self._busy:
            return
        paths, _ = QFileDialog.getOpenFileNames(
            self, "파일 열기", self.last_path,
            "지원 파일 (*.pdf *.txt *.jpg *.jpeg *.png *.bmp *.docx *.pptx *.xlsx)")
        if paths:
            self.load_files(paths, clear_before=clear)

    def _convert_office_to_pdf(self, input_path: str, output_path: str) -> None:
        if sys.platform != "win32":
            raise RuntimeError("MS Office 불러오기는 Windows에서만 지원합니다.")
        import comtypes
        import comtypes.client
        app: Any = None
        document: Any = None
        ext = os.path.splitext(input_path)[1].lower()
        comtypes.CoInitialize()
        try:
            if ext == ".docx":
                app = comtypes.client.CreateObject("Word.Application")
                app.Visible = False
                app.DisplayAlerts = 0
                app.AutomationSecurity = 3
                document = app.Documents.Open(
                    os.path.abspath(input_path), ReadOnly=True,
                    AddToRecentFiles=False)
                document.ExportAsFixedFormat(
                    OutputFileName=output_path, ExportFormat=17, OptimizeFor=1)
            elif ext == ".pptx":
                app = comtypes.client.CreateObject("PowerPoint.Application")
                app.AutomationSecurity = 3
                document = app.Presentations.Open(
                    os.path.abspath(input_path), ReadOnly=True, WithWindow=False)
                document.ExportAsFixedFormat(output_path, 2)
            else:
                app = comtypes.client.CreateObject("Excel.Application")
                app.Visible = False
                app.DisplayAlerts = False
                app.AutomationSecurity = 3
                document = app.Workbooks.Open(
                    os.path.abspath(input_path), UpdateLinks=0, ReadOnly=True)
                document.ExportAsFixedFormat(0, output_path, Quality=1)
        finally:
            try:
                if document is not None:
                    if ext == ".pptx":
                        document.Close()
                    else:
                        document.Close(False)
            finally:
                try:
                    if app is not None:
                        app.Quit()
                finally:
                    comtypes.CoUninitialize()

    def open_source_document(self, path: str) -> SourceDoc:
        ext = os.path.splitext(path)[1].lower().lstrip(".")
        temps: list[str] = []
        render = path
        doc: Optional[fitz.Document] = None
        try:
            kind = "pdf"
            if ext != "pdf":
                fd, render = tempfile.mkstemp(suffix=".pdf", prefix="pdf_master_")
                os.close(fd)
                temps.append(render)
                if ext in {"docx", "pptx", "xlsx"}:
                    kind = "office"
                    # Office 호출을 별도 스레드에서 수행합니다.
                    # PDF 렌더 작업과는 동시에 실행하지 않습니다.
                    import concurrent.futures
                    with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
                        future = pool.submit(self._convert_office_to_pdf, path, render)
                        while not future.done():
                            QApplication.processEvents()
                            import time
                            time.sleep(0.02)
                        future.result()
                elif ext == "txt":
                    kind = "txt"
                    with self.create_novel_pdf_from_text(read_text(path)) as converted:
                        converted.save(render)
                else:
                    kind = "image"
                    with fitz.open(path) as image:
                        pdf_bytes = image.convert_to_pdf()
                    with open(render, "wb") as output:
                        output.write(pdf_bytes)
            doc = fitz.open(render)
            if doc.needs_pass:
                raise ValueError("암호화된 PDF입니다. 암호를 해제한 파일을 사용하세요.")
            if not doc.page_count:
                raise ValueError("페이지가 없는 문서입니다.")
            return SourceDoc(doc, os.path.basename(path), path, kind, render, temps)
        except Exception:
            if doc is not None:
                doc.close()
            for temp in temps:
                try:
                    os.remove(temp)
                except OSError:
                    pass
            raise

    def load_files(self, files: list[str], clear_before: bool = False,
                   source: str = "") -> None:
        if self._busy:
            return
        paths = [p for p in files if os.path.isfile(p) and
                 os.path.splitext(p)[1].lower().lstrip(".") in self.SUPPORTED_EXTS]
        if not paths:
            return
        # TXT 다이얼로그는 목록을 재귀적으로 변경하지 않고 결과 경로만 반환합니다.
        final: list[str] = []
        for path in paths:
            if path.lower().endswith(".txt"):
                try:
                    dialog = TextImportDialog(self, path)
                    dialog.exec()
                    if dialog.saved_pdf:
                        final.append(dialog.saved_pdf)
                except Exception as exc:
                    QMessageBox.warning(self, "텍스트 읽기 오류", str(exc))
            else:
                final.append(path)
        if not final:
            return
        if clear_before:
            self.clear_all()
        if not self._begin(len(final), "파일 로드"):
            return
        try:
            for number, path in enumerate(final, 1):
                try:
                    self.set_status("파일 처리 중: " + os.path.basename(path))
                    sd = self.open_source_document(path)
                    d_id = len(self.docs)
                    self.docs.append(sd)
                    for index in range(sd.doc.page_count):
                        item = QListWidgetItem(f"[{sd.name}] P.{index + 1}")
                        item.setData(ROLE, (d_id, index))
                        self.orig_list.addItem(item)
                        self._orig_item_map[(d_id, index)] = item
                except Exception as exc:
                    QMessageBox.warning(self, "파일 로드 오류",
                                        os.path.basename(path) + "\n\n" + str(exc))
                self._advance(number)
            self.update_last_path(final[0])
            self._update_list_icon_size()
            self.set_status(f"로드 완료: {len(self.docs)}개 문서 / "
                            f"{self.orig_list.count()}개 페이지")
        finally:
            self._end()
        for d_id, index in list(self._orig_item_map)[-self.THUMB_QUEUE_LIMIT_PER_LOAD:]:
            self._enqueue_thumbnail_task(d_id, index)
        self._apply_filter(self.orig_list, self.orig_search.text())

    def create_novel_pdf_from_text(self, text: str) -> fitz.Document:
        doc = fitz.open()
        dimensions = {
            "A5": (148, 210), "신국판": (152, 225),
            "46판": (128, 188), "B5": (182, 257)}
        mm_w, mm_h = dimensions.get(self.novel_size_mode, (148, 210))
        width, height = mm_w * 72 / 25.4, mm_h * 72 / 25.4
        margin = width * 0.107
        top, bottom = height * 0.085, height * 0.095
        font_size = int(self.text_import_font_size)
        line_height = font_size * 1.75
        try:
            font = fitz.Font(fontfile=self.font_path) if self.font_path else fitz.Font("korea")
        except Exception:
            font = fitz.Font("korea")
        alias = "KOR" if self.font_path else "korea"
        # 실제 측정에 사용한 폰트와 삽입 폰트를 동일하게 유지합니다.
        if not self.font_path:
            alias = "korea"

        def new_page() -> fitz.Page:
            page = doc.new_page(width=width, height=height)
            if self.font_path:
                page.insert_font(fontname=alias, fontfile=self.font_path)
            return page

        try:
            page = new_page()
            y = top + font_size
            for paragraph in text.splitlines():
                paragraph = paragraph.strip()
                if not paragraph:
                    y += line_height * 0.5
                    continue
                lines: list[str] = []
                current = ""
                current_width = 0.0
                for char in paragraph:
                    char_width = font.text_length(char, fontsize=font_size)
                    limit = width - 2 * margin - (font_size if not lines else 0)
                    if current and current_width + char_width > limit:
                        lines.append(current)
                        current = ""
                        current_width = 0.0
                    current += char
                    current_width += char_width
                if current:
                    lines.append(current)
                for number, line in enumerate(lines):
                    if y > height - bottom:
                        page = new_page()
                        y = top + font_size
                    x = margin + (font_size if number == 0 else 0)
                    page.insert_text((x, y), line, fontsize=font_size, fontname=alias)
                    y += line_height
            for index in range(doc.page_count):
                doc[index].insert_textbox(
                    fitz.Rect(margin, height - bottom + 15, width - margin, height - 10),
                    f"- {index + 1} -", fontsize=9, fontname="helv", align=1)
            return doc
        except Exception:
            doc.close()
            raise

    def _enqueue_thumbnail_task(self, d_id: int, index: int) -> None:
        if self._busy or self._closing:
            return
        key = ("thumb", self._generation, self._thumb_generation, d_id, index)
        if key in self._inflight:
            return
        self._inflight.add(key)
        self.thread_pool.start(RenderTask(self.render_signals, {
            "kind": "thumb", "key": key, "generation": self._generation,
            "thumb_generation": self._thumb_generation,
            "path": self.docs[d_id].render_path, "index": index,
            "d_id": d_id, "width": self._icon_px}))

    def _effective_rotation(self, d_id: int, index: int, rotation: int) -> int:
        rect = self.docs[d_id].doc[index].rect
        w, h = (rect.height, rect.width) if rotation % 180 else (rect.width, rect.height)
        if ((self.merge_rotate_mode == "portrait" and w > h) or
                (self.merge_rotate_mode == "landscape" and h > w)):
            rotation += 90
        return rotation % 360

    def _page_layout(self, key: PageKey) -> tuple[int, tuple[float, float], bool]:
        d_id, index, rotation = key
        rotation = self._effective_rotation(d_id, index, rotation)
        rect = self.docs[d_id].doc[index].rect
        w, h = (rect.height, rect.width) if rotation % 180 else (rect.width, rect.height)
        if self.merge_size_mode == "a4":
            w, h = 210 * 72 / 25.4, 297 * 72 / 25.4
        elif self.merge_size_mode == "first":
            first = self.work_list.item(0)
            data = first.data(ROLE) if first is not None else key
            d0, i0 = data[:2]
            r0 = self._effective_rotation(d0, i0, data[2] if len(data) > 2 else 0)
            rect0 = self.docs[d0].doc[i0].rect
            w, h = ((rect0.height, rect0.width) if r0 % 180
                    else (rect0.width, rect0.height))
        if ((self.merge_orient_mode == "portrait" and w > h) or
                (self.merge_orient_mode == "landscape" and h > w)):
            w, h = h, w
        margins = self.merge_size_mode != "keep" or self.merge_orient_mode != "keep"
        return rotation, (float(w), float(h)), margins

    def handle_item_preview(self, item: Optional[QListWidgetItem]) -> None:
        if item is None or self._busy:
            return
        data = item.data(ROLE)
        if not data:
            return
        d_id, index = data[:2]
        self._pending_preview_key = (d_id, index, data[2] if len(data) > 2 else 0)
        self._preview_timer.start(self.PREVIEW_DEBOUNCE_MS)
        if (d_id, index) not in self._thumb_cache:
            self._enqueue_thumbnail_task(d_id, index)

    def _render_pending_preview(self) -> None:
        key = self._pending_preview_key
        if key is None or self._busy or self._closing:
            return
        d_id, index, _ = key
        if d_id >= len(self.docs):
            return
        rotation, size, margins = self._page_layout(key)
        cache_key = (
            "preview", self._generation, d_id, index, rotation, size,
            margins, self.merge_align_mode)
        self._current_preview_key = cache_key
        cached = self._preview_cache.get(cache_key)
        if cached is not None:
            self._preview_cache.move_to_end(cache_key)
            self.preview_label.set_pixmap_fixed(cached)
            return
        if cache_key in self._inflight:
            return
        self._inflight.add(cache_key)
        self.thread_pool.start(RenderTask(self.render_signals, {
            "kind": "preview", "key": cache_key, "generation": self._generation,
            "path": self.docs[d_id].render_path, "index": index,
            "rotation": rotation, "size": size, "margins": margins,
            "align": self.merge_align_mode, "width": self.PREVIEW_MAX_W}))

    @Slot(object)
    def _on_render_ready(self, out: dict[str, Any]) -> None:
        self._inflight.discard(out["key"])
        if self._closing or out["generation"] != self._generation:
            return
        image = QImage(out["samples"], out["w"], out["h"], out["stride"],
                       QImage.Format.Format_RGB888).copy()
        pixmap = QPixmap.fromImage(image)
        if out["kind"] == "thumb":
            if out["thumb_generation"] != self._thumb_generation:
                return
            key = (out["d_id"], out["index"])
            self._thumb_cache[key] = pixmap
            self._thumb_cache.move_to_end(key)
            if len(self._thumb_cache) > self.THUMB_CACHE_LIMIT:
                self._thumb_cache.popitem(last=False)
            item = self._orig_item_map.get(key)
            if item is not None:
                item.setIcon(QIcon(pixmap))
            for index in range(self.work_list.count()):
                item = self.work_list.item(index)
                if item is not None and tuple(item.data(ROLE)[:2]) == key:
                    item.setIcon(QIcon(pixmap))
        else:
            self._preview_cache[out["key"]] = pixmap
            self._preview_cache.move_to_end(out["key"])
            if len(self._preview_cache) > self.PREVIEW_CACHE_LIMIT:
                self._preview_cache.popitem(last=False)
            if self._current_preview_key == out["key"]:
                self.preview_label.set_pixmap_fixed(pixmap)

    @Slot(object)
    def _on_render_error(self, out: dict[str, Any]) -> None:
        self._inflight.discard(out["key"])
        if out["generation"] == self._generation and not self._closing:
            self.set_status("렌더 오류: " + out["message"])

    def refresh_current_preview(self) -> None:
        self._preview_cache.clear()
        item = self.work_list.currentItem()
        if item is None:
            item = self.orig_list.currentItem()
        if item is not None:
            self.handle_item_preview(item)
        elif not self.orig_list.count():
            self._pending_preview_key = None
            self._current_preview_key = None
            self.preview_label.clear()

    def on_worklist_reordered(self) -> None:
        self.refresh_current_preview()
        self.set_status("작업 리스트 순서 변경됨")

    def add_selected_to_work(self, item=None) -> None:
        if self._busy:
            return
        count = 0
        for source in self.orig_list.selectedItems():
            d_id, index = source.data(ROLE)
            key = (d_id, index)
            if key in self._work_key_set:
                continue
            target = QListWidgetItem(source.text())
            target.setData(ROLE, (d_id, index, 0))
            target.setIcon(source.icon())
            self.work_list.addItem(target)
            self._work_key_set.add(key)
            count += 1
        self._update_list_icon_size()
        self._apply_filter(self.work_list, self.work_search.text())
        self.refresh_current_preview()
        self.set_status(f"{count}개 추가 (중복 제외)")

    def remove_work_items(self, item=None) -> None:
        if self._busy:
            return
        selected = self.work_list.selectedItems()
        for target in selected:
            self._work_key_set.discard(tuple(target.data(ROLE)[:2]))
            self.work_list.takeItem(self.work_list.row(target))
        self.refresh_current_preview()
        self.set_status(f"{len(selected)}개 삭제")

    def rotate_work_items(self, angle: int, reset: bool = False) -> None:
        if self._busy:
            return
        for item in self.work_list.selectedItems():
            data = item.data(ROLE)
            d_id, index = data[:2]
            rotation = 0 if reset else ((data[2] if len(data) > 2 else 0) + angle) % 360
            item.setData(ROLE, (d_id, index, rotation))
            title = re.sub(r"\s+\[회전:\s*\d+°\]$", "", item.text())
            item.setText(title + (f" [회전: {rotation}°]" if rotation else ""))
        self.refresh_current_preview()

    def clear_all(self) -> None:
        if self._busy:
            return
        self._preview_timer.stop()
        self._generation += 1
        self.thread_pool.clear()
        self.thread_pool.waitForDone()
        for sd in self.docs:
            sd.close()
            sd.cleanup()
        self.docs.clear()
        self._preview_cache.clear()
        self._thumb_cache.clear()
        self._orig_item_map.clear()
        self._inflight.clear()
        self._work_key_set.clear()
        self._pending_preview_key = None
        self._current_preview_key = None
        self.orig_list.clear()
        self.work_list.clear()
        self.orig_search.clear()
        self.work_search.clear()
        self.preview_label.clear()
        self.set_status("초기화 완료")

    def _work_items(self) -> list[PageKey]:
        result: list[PageKey] = []
        for index in range(self.work_list.count()):
            item = self.work_list.item(index)
            if item is not None:
                data = item.data(ROLE)
                result.append((data[0], data[1], data[2] if len(data) > 2 else 0))
        return result

    def extract_text(self, page: fitz.Page) -> str:
        result = page.get_text("text", sort=self.text_mode == "layout")
        return result.strip() if isinstance(result, str) else ""

    def render_page_image_bytes(self, page: fitz.Page, rotation: int = 0,
                                fmt: str = "png") -> io.BytesIO:
        fmt = "jpeg" if fmt.lower() in {"jpg", "jpeg"} else fmt.lower()
        matrix = fitz.Matrix(self.current_dpi / 72, self.current_dpi / 72).prerotate(rotation)
        pix = page.get_pixmap(matrix=matrix, colorspace=fitz.csRGB, alpha=False)
        raw = (pix.tobytes("jpeg", jpg_quality=self._get_jpg_quality())
               if fmt == "jpeg" else pix.tobytes(fmt))
        return io.BytesIO(raw)

    def action_merge(self) -> None:
        items = self._work_items()
        if not items or self._busy:
            return
        path = self.pick_save("병합완료_결과물", "pdf")
        if not path or not self._begin(len(items), "병합"):
            return
        try:
            with fitz.open() as out:
                for number, key in enumerate(items, 1):
                    d_id, index, _ = key
                    rotation, size, margins = self._page_layout(key)
                    if not margins and not self.pdf_rasterize_mode:
                        with fitz.open() as temp:
                            temp.insert_pdf(self.docs[d_id].doc,
                                            from_page=index, to_page=index)
                            page = temp[0]
                            page.set_rotation((page.rotation + rotation) % 360)
                            out.insert_pdf(temp)
                    else:
                        build_page(out, self.docs[d_id].doc, index, rotation,
                                   size, self.merge_align_mode, margins,
                                   self.pdf_rasterize_mode, self.current_dpi,
                                   self._get_jpg_quality())
                    self._advance(number)
                self._save_pdf(out, path)
            self.set_status("PDF 병합 완료")
        except Exception as exc:
            QMessageBox.warning(self, "PDF 병합 오류", str(exc))
        finally:
            self._end()

    def action_extract(self) -> None:
        items = self._work_items()
        if not items or self._busy:
            return
        path = self.pick_save("추출페이지", "pdf")
        if not path:
            return
        base = os.path.splitext(path)[0]
        existing = [f"{base}_{i + 1}.pdf" for i in range(len(items))
                    if os.path.exists(f"{base}_{i + 1}.pdf")]
        if existing and QMessageBox.question(
                self, "덮어쓰기", f"기존 파일 {len(existing)}개를 덮어쓸까요?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.No) != QMessageBox.StandardButton.Yes:
            return
        if not self._begin(len(items), "추출"):
            return
        try:
            for number, (d_id, index, rotation) in enumerate(items, 1):
                with fitz.open() as out:
                    source = self.docs[d_id].doc
                    if self.pdf_rasterize_mode:
                        rect = source[index].rect
                        size = ((rect.height, rect.width) if rotation % 180
                                else (rect.width, rect.height))
                        build_page(out, source, index, rotation, size, "center",
                                   False, True, self.current_dpi,
                                   self._get_jpg_quality())
                    else:
                        out.insert_pdf(source, from_page=index, to_page=index)
                        page = out[0]
                        page.set_rotation((page.rotation + rotation) % 360)
                    self._save_pdf(out, f"{base}_{number}.pdf")
                self._advance(number)
            self.set_status("개별 추출 완료")
        except Exception as exc:
            QMessageBox.warning(self, "PDF 추출 오류", str(exc))
        finally:
            self._end()

    def _run_conversion(self, extension: str, callback) -> None:
        items = self._work_items()
        if not items or self._busy:
            return
        path = self.pick_save("변환결과", extension)
        if not path or not self._begin(len(items), extension.upper() + " 변환"):
            return
        try:
            callback(path, items)
            self.set_status(extension.upper() + " 변환 완료")
        except ImportError as exc:
            QMessageBox.warning(self, "패키지 필요", str(exc))
        except Exception as exc:
            QMessageBox.warning(self, "변환 오류", str(exc))
        finally:
            self._end()

    def convert_to_word(self) -> None:
        self._run_conversion("docx", self._write_word)

    def _write_word(self, path: str, items: list[PageKey]) -> None:
        from docx import Document
        from docx.shared import Inches
        doc = Document()
        for number, (d_id, index, rotation) in enumerate(items, 1):
            page = self.docs[d_id].doc[index]
            text = self.extract_text(page)
            if self.convert_mode in {"image", "mixed"} or not text:
                available = Inches(6)
                from PIL import Image
                stream = self.render_page_image_bytes(page, rotation, "jpg")
                with Image.open(stream) as image:
                    iw, ih = image.size
                stream.seek(0)
                max_height = int(Inches(8.5))
                scaled_width = min(int(available), int(max_height * iw / ih))
                doc.add_picture(stream, width=Inches(scaled_width / 914400))
            if self.convert_mode != "image" and text:
                if self.convert_mode == "mixed":
                    doc.add_paragraph("[추출 텍스트]")
                doc.add_paragraph(text)
            if number < len(items):
                doc.add_page_break()
            self._advance(number)
        doc.save(path)

    def convert_to_ppt(self) -> None:
        self._run_conversion("pptx", self._write_ppt)

    def _write_ppt(self, path: str, items: list[PageKey]) -> None:
        from pptx import Presentation
        from pptx.util import Pt, Emu
        from pptx.enum.text import MSO_AUTO_SIZE
        from PIL import Image
        prs = Presentation()
        d0, i0, r0 = items[0]
        rect = self.docs[d0].doc[i0].rect
        first_w, first_h = ((rect.height, rect.width) if r0 % 180
                            else (rect.width, rect.height))
        width, height = Pt(first_w), Pt(first_h)
        prs.slide_width, prs.slide_height = width, height
        sw, sh = int(width), int(height)
        blank = next((layout for layout in prs.slide_layouts
                      if not len(layout.placeholders)), prs.slide_layouts[6])
        for number, (d_id, index, rotation) in enumerate(items, 1):
            page = self.docs[d_id].doc[index]
            text = self.extract_text(page)
            slide = prs.slides.add_slide(blank)
            if self.convert_mode in {"image", "mixed"} or not text:
                stream = self.render_page_image_bytes(page, rotation)
                with Image.open(stream) as image:
                    iw, ih = image.size
                stream.seek(0)
                scale = min(sw / iw, sh / ih)
                w, h = int(iw * scale), int(ih * scale)
                slide.shapes.add_picture(
                    stream, Emu((sw - w) // 2), Emu((sh - h) // 2),
                    width=Emu(w), height=Emu(h))
            if self.convert_mode != "image" and text:
                if self.convert_mode == "mixed":
                    # 혼합 모드는 이미지 위에 본문을 겹치지 않고 발표자 노트에 보관.
                    notes_frame = slide.notes_slide.notes_text_frame
                    if notes_frame is not None:
                        notes_frame.text = text
                else:
                    box = slide.shapes.add_textbox(Emu(0), Emu(0),
                                                   Emu(sw), Emu(sh))
                    tf = box.text_frame
                    tf.word_wrap = True
                    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                    tf.text = text
                    for paragraph in tf.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "Malgun Gothic"
                            run.font.size = Pt(20)
            self._advance(number)
        prs.save(path)

    def convert_to_excel(self) -> None:
        self._run_conversion("xlsx", self._write_excel)

    def _write_excel(self, path: str, items: list[PageKey]) -> None:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        wb = Workbook()
        wb.remove(wb.worksheets[0])
        streams: list[io.BytesIO] = []
        try:
            for number, (d_id, index, rotation) in enumerate(items, 1):
                page = self.docs[d_id].doc[index]
                text = self.extract_text(page)
                ws = wb.create_sheet(f"Page_{number}")
                ws.column_dimensions["A"].width = 100
                row = 1
                if self.convert_mode in {"image", "mixed"} or not text:
                    stream = self.render_page_image_bytes(page, rotation)
                    streams.append(stream)
                    image = XLImage(stream)
                    ws.add_image(image, "A1")
                    row = int(image.height / 20) + 3
                if self.convert_mode != "image" and text:
                    for line in text.splitlines():
                        cell = ws.cell(row=row, column=1)
                        cell.value = ILLEGAL_CHARACTERS_RE.sub("", line)[:32767]
                        cell.data_type = "s"  # '='로 시작하는 추출문도 수식으로 실행하지 않음
                        row += 1
                self._advance(number)
            wb.save(path)
        finally:
            wb.close()
            for stream in streams:
                stream.close()

    def convert_to_text(self) -> None:
        self._run_conversion("txt", self._write_text)

    def _write_text(self, path: str, items: list[PageKey]) -> None:
        with open(path, "w", encoding="utf-8") as output:
            for number, (d_id, index, _) in enumerate(items, 1):
                text = self.extract_text(self.docs[d_id].doc[index])
                output.write(f"--- Page {number} ---\n{text}\n\n")
                self._advance(number)

    def convert_to_epub(self) -> None:
        self._run_conversion("epub", self._write_epub)

    def _write_epub(self, path: str, items: list[PageKey]) -> None:
        from ebooklib import epub
        book = epub.EpubBook()
        book.set_identifier(uuid.uuid4().hex)
        book.set_title("PDF 변환 문서")
        book.set_language("ko")
        chapters: list[Any] = []
        style = epub.EpubItem(
            uid="style", file_name="style/book.css", media_type="text/css",
            content=b"body{font-family:sans-serif}p{line-height:1.7}"
                    b"img{max-width:100%;height:auto}")
        book.add_item(style)
        for number, (d_id, index, rotation) in enumerate(items, 1):
            page = self.docs[d_id].doc[index]
            text = self.extract_text(page)
            chapter = epub.EpubHtml(title=f"Page {number}",
                                   file_name=f"page_{number}.xhtml", lang="ko")
            parts = [f"<h2>Page {number}</h2>"]
            if self.convert_mode in {"image", "mixed"} or not text:
                name = f"images/page_{number}.jpg"
                image = epub.EpubItem(
                    uid=f"image_{number}", file_name=name, media_type="image/jpeg",
                    content=self.render_page_image_bytes(page, rotation, "jpg").getvalue())
                book.add_item(image)
                parts.append(f'<p><img src="{name}" alt="Page {number}"/></p>')
            if self.convert_mode != "image" and text:
                for paragraph in text.split("\n\n"):
                    escaped = html.escape(paragraph).replace("\n", "<br/>")
                    parts.append("<p>" + escaped + "</p>")
            chapter.content = "".join(parts)
            chapter.add_item(style)
            book.add_item(chapter)
            chapters.append(chapter)
            self._advance(number)
        book.toc = list(chapters)
        book.spine = ["nav"] + chapters
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(path, book, {})

    def convert_to_images(self) -> None:
        items = self._work_items()
        if not items or self._busy:
            return
        path, selected = QFileDialog.getSaveFileName(
            self, "이미지 저장", self.default_save_path("이미지변환", "png"),
            "PNG (*.png);;JPEG (*.jpg *.jpeg);;GIF (*.gif);;BMP (*.bmp);;TIFF (*.tiff *.tif)")
        if not path:
            return
        self.update_last_path(path)
        ext = next((value for label, value in (
            ("JPEG", "jpg"), ("GIF", "gif"), ("BMP", "bmp"), ("TIFF", "tiff"))
                    if label in selected), "png")
        base = os.path.splitext(path)[0]
        existing = [f"{base}_{i + 1:03d}.{ext}" for i in range(len(items))
                    if os.path.exists(f"{base}_{i + 1:03d}.{ext}")]
        if existing and QMessageBox.question(
                self, "덮어쓰기", f"기존 이미지 {len(existing)}개를 덮어쓸까요?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.No) != QMessageBox.StandardButton.Yes:
            return
        if not self._begin(len(items), "이미지 변환"):
            return
        try:
            from PIL import Image
            formats = {"png": "PNG", "jpg": "JPEG", "gif": "GIF",
                       "bmp": "BMP", "tiff": "TIFF"}
            for number, (d_id, index, rotation) in enumerate(items, 1):
                stream = self.render_page_image_bytes(
                    self.docs[d_id].doc[index], rotation)
                with Image.open(stream) as image:
                    options: dict[str, Any] = {}
                    if ext != "gif":
                        options["dpi"] = (self.current_dpi, self.current_dpi)
                    if ext == "jpg":
                        options["quality"] = self._get_jpg_quality()
                        image = image.convert("RGB")
                    image.save(f"{base}_{number:03d}.{ext}",
                               format=formats[ext], **options)
                self._advance(number)
            self.set_status("이미지 변환 완료")
        except Exception as exc:
            QMessageBox.warning(self, "이미지 변환 오류", str(exc))
        finally:
            self._end()

    def load_settings(self) -> None:
        geometry = self.settings.value("geometry")
        if isinstance(geometry, QByteArray):
            self.restoreGeometry(geometry)
        splitter = self.settings.value("splitter_state")
        if isinstance(splitter, QByteArray):
            self.content_splitter.restoreState(splitter)
        else:
            self.content_splitter.setSizes([350, 350, 700])
        self.change_ui_font(self.current_font_size)

    def closeEvent(self, event):
        if self._busy:
            QMessageBox.information(self, "작업 중", "현재 작업이 끝난 뒤 종료하세요.")
            event.ignore()
            return
        self._closing = True
        self._preview_timer.stop()
        self._generation += 1
        self.thread_pool.clear()
        self.thread_pool.waitForDone()
        for sd in self.docs:
            sd.close()
            sd.cleanup()
        self.settings.setValue("geometry", self.saveGeometry())
        self.settings.setValue("splitter_state", self.content_splitter.saveState())
        for key, value in (
            ("last_path", self.last_path), ("font_size", self.current_font_size),
            ("dpi", self.current_dpi),
            ("text_import_font_size", self.text_import_font_size),
            ("text_mode", self.text_mode), ("convert_mode", self.convert_mode),
            ("novel_size_mode", self.novel_size_mode),
            ("merge_size_mode", self.merge_size_mode),
            ("merge_align_mode", self.merge_align_mode),
            ("merge_rotate_mode", self.merge_rotate_mode),
            ("merge_orient_mode", self.merge_orient_mode),
            ("pdf_save_mode", self.pdf_save_mode),
            ("pdf_rasterize_mode", self.pdf_rasterize_mode),
        ):
            self.settings.setValue(key, value)
        self.settings.sync()
        event.accept()


if __name__ == "__main__":
    QLoggingCategory.setFilterRules("qt.gui.icc.warning=false")

    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    win = IntegratedPDFMaster()
    win.show()
    sys.exit(app.exec())
