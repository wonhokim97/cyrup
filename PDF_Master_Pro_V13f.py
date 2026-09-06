# -*- coding: utf-8 -*-
"""PDF Master Pro (ALL-IN-ONE) - Final Build (Patched Integrated)

✅ 포함 기능
- 파일 열기(새로 열기/추가), 원본 리스트 드래그&드롭으로 파일 추가
- 원본 리스트 ↔ 작업 리스트(다중 선택), 작업 리스트 내부 드래그 재정렬
- 실시간 미리보기(hover/클릭/휠), 썸네일/프리뷰 백그라운드 렌더링(QThreadPool)
- 썸네일/프리뷰 캐시(LRU)
- PDF 병합(페이지 크기/배치/자동 회전 옵션), 작업 리스트 개별 추출
- 변환: Word / PPT / Excel / Text / Images
- 설정 저장/복원(QSettings): last_path, DPI, UI 폰트 크기, TXT 임포트 폰트 크기,
  텍스트 추출 모드, 변환 모드, 병합 옵션, splitter 상태, 창 크기, PDF 저장 모드

✅ 누적 패치 반영(요청 반영)
1) PDF 저장 모드(일반/저용량/고품질)
2) 리스트 아이템 간격 최소화
3) 이미지 렌더링 Matrix+prerotate
4) 미리보기는 항상 캐시 PDF(render_path) 기반(빠른 렌더)
5) PDF 캐시 중복 생성 방지(원본 경로 + mtime + size)
6) TXT → PDF 변환은 fpdf2 조판(uni=True 폰트 등록 포함) 후 저장→재오픈
7) 병합 결과에 페이지 번호 삽입: 홀수(오른쪽) / 짝수(왼쪽)  "- 현재/전체 -"
   - insert_textbox() 사용으로 정렬/버전 호환성 강화
8) ✅ 실시간 미리보기 고해상도 개선:
   - 프리뷰 위젯 폭 × DevicePixelRatio 기반 목표 렌더 폭 산정
   - 저장 모드에 따라 품질 계수/상한 자동 조절
   - 캐시에는 (pixmap, rendered_w) 저장 → 더 큰 해상도 필요 시 자동 재렌더

필수 패키지:
  pip install PySide6 PyMuPDF python-docx openpyxl python-pptx fpdf2

권장 폰트(동봉 권장):
- ./fonts/NotoSansKR-Regular.ttf
- (권장) ./fonts/NotoSansKR-Bold.ttf
"""

import sys
import os
import io
import re
import datetime
import tempfile
import shutil
from dataclasses import dataclass, field
from collections import OrderedDict
from typing import Optional, Tuple, List, Dict, Any, Set

import fitz  # PyMuPDF

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QFileDialog, QVBoxLayout, QHBoxLayout, QWidget,
    QPushButton, QListWidget, QListWidgetItem, QLabel, QScrollArea, QMenu,
    QAbstractItemView, QSplitter, QFrame, QMessageBox, QProgressBar, QLineEdit,
    QSizePolicy
)
from PySide6.QtGui import (
    QPixmap, QAction, QFont, QCursor, QActionGroup, QIcon, QFontMetrics, QFontDatabase
)
from PySide6.QtCore import (
    Qt, QSettings, QTimer, QObject, Signal, QRunnable, QThreadPool, Slot, QSize,
    QCoreApplication
)

from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from pptx import Presentation

# --- fpdf2 (TXT -> PDF 조판) ---
try:
    from fpdf import FPDF
except Exception:
    FPDF = None


# ============================================================
# Utils
# ============================================================

def resource_base_dir() -> str:
    base = getattr(sys, "_MEIPASS", None)
    if base and os.path.isdir(base):
        return base
    return os.path.dirname(os.path.abspath(__file__))


# ============================================================
# Text typesetting helpers (참조 코드 핵심만 이식)
# ============================================================

CHAPTER_PATTERNS = [
    r'^\s*#{1,6}\s+(.+?)\s*$',
    r'^\s*(제\s*\d+\s*(화|장)|\d+\s*(화|장))\s*[:.\-]?\s*(.*)\s*$',
    r'^\s*(Chapter|CHAPTER|Episode|EPISODE)\s*\d+\s*[:.\-]?\s*(.*)\s*$',
    r'^\s*(프롤로그|에필로그|외전|후기|서장|서문|머리말|끝)\s*[:.\-]?\s*(.*)\s*$',
]
DIALOGUE_PREFIXES = ('"', "'", '“', '‘', '「', '『', '—', '―', '-', '–', '…')
SCENE_BREAK_RE = re.compile(r'^[\s\*\-_=~·•\.]{3,}$')
LISTLIKE_RE = re.compile(r'^\s*(?:\d+[\.\)]|[※▶■◆•\-])\s+')
SENTENCE_END_RE = re.compile(r'.*[\.!?…]["”’\')\]]*\s*$')
KOR_END_RE = re.compile(r'.*(다\.|요\.|죠\.|네\.|임\.|함\.|다!|다\?|요!|요\?)\s*$')


def detect_chapter_title(line: str, max_title_len: int = 36):
    for pat in CHAPTER_PATTERNS:
        m = re.match(pat, line)
        if m:
            groups = [g for g in m.groups() if g is not None]
            if pat.startswith(r'^\s*#{'):
                t = groups[0].strip()
                return t if len(t) <= max_title_len else None
            title_main = (groups[0].strip() if groups else line.strip())
            subtitle = groups[-1].strip() if len(groups) >= 2 else ""
            t = (title_main + (f" - {subtitle}" if subtitle and subtitle != title_main else "")).strip()
            return t if len(t) <= max_title_len else None
    return None


def _count_nonempty(lines: List[str]) -> int:
    return sum(1 for s in lines if s.strip())


def split_into_chapters(text: str, strict: bool = True, min_body_lines: int = 3):
    """
    ✅ 오탐 감소 패치:
    - 직전이 공백행(또는 문서 시작)일 때만 장 제목 판정 허용
    """
    lines = text.replace('\r\n', '\n').replace('\r', '\n').split('\n')
    chapters = []
    cur_title = "본문"
    cur_lines: List[str] = []

    prev_blank = True  # 문서 시작은 장 후보 허용

    for line in lines:
        t = detect_chapter_title(line) if prev_blank else None

        if t:
            if strict and chapters and _count_nonempty(cur_lines) < min_body_lines:
                cur_lines.append(line)
            else:
                if cur_lines and any(s.strip() for s in cur_lines):
                    chapters.append({"title": cur_title, "lines": cur_lines})
                cur_title = t
                cur_lines = []
        else:
            cur_lines.append(line)

        prev_blank = (not line.strip())

    if cur_lines and any(s.strip() for s in cur_lines):
        chapters.append({"title": cur_title, "lines": cur_lines})

    return chapters if chapters else [{"title": "본문", "lines": lines}]


def is_scene_break(line: str) -> bool:
    return bool(SCENE_BREAK_RE.match(line.strip()))


def is_dialogue(line: str) -> bool:
    s = line.lstrip()
    return s.startswith(DIALOGUE_PREFIXES)


def looks_like_list(line: str) -> bool:
    return bool(LISTLIKE_RE.match(line))


def smart_make_paragraphs(lines: List[str]) -> List[str]:
    paras: List[str] = []
    buf: List[str] = []
    buf_len = 0

    def flush():
        nonlocal buf_len
        if buf:
            paras.append(" ".join(buf).strip())
            buf.clear()
            buf_len = 0

    for raw in lines:
        ln = raw.rstrip("\n")
        s = ln.strip()

        if not s:
            flush()
            if not paras or paras[-1] != "":
                paras.append("")
            continue

        if is_scene_break(s):
            flush()
            if not paras or paras[-1] != "":
                paras.append("")
            continue

        if looks_like_list(s):
            flush()
            paras.append(s)
            continue

        if is_dialogue(s):
            flush()
            paras.append(s)
            continue

        buf.append(s)
        buf_len += len(s)
        if buf_len >= 240 and (SENTENCE_END_RE.match(s) or KOR_END_RE.match(s)):
            flush()

    flush()
    while paras and paras[-1] == "":
        paras.pop()
    return paras


def contains_korean(text: str) -> bool:
    for ch in text:
        o = ord(ch)
        if (0xAC00 <= o <= 0xD7A3) or (0x1100 <= o <= 0x11FF) or (0x3130 <= o <= 0x318F):
            return True
    return False


# ============================================================
# Data Model
# ============================================================

@dataclass
class SourceDoc:
    doc: fitz.Document
    name: str
    path: str
    source_type: str  # 'pdf' / 'txt' / 'image'
    render_path: str  # thread-safe reopen target (PDF path)
    temp_paths: List[str] = field(default_factory=list)

    def close(self):
        try:
            self.doc.close()
        except Exception:
            pass

    def cleanup(self):
        for p in list(self.temp_paths):
            try:
                if p and os.path.exists(p):
                    os.remove(p)
            except Exception:
                pass
        self.temp_paths.clear()


# ============================================================
# Responsive Preview Label
# ============================================================

class ResponsiveLabel(QLabel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAlignment(Qt.AlignCenter)
        self._current_pixmap: Optional[QPixmap] = None
        self.setMinimumWidth(0)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.setText("파일을 로드하고 리스트에 마우스를 올리세요")

    def set_pixmap_fixed(self, pixmap: QPixmap):
        self._current_pixmap = pixmap
        self.update_display()

    def clear(self):
        self._current_pixmap = None
        super().clear()
        self.setText("파일을 로드하고 리스트에 마우스를 올리세요")

    def update_display(self):
        if self._current_pixmap:
            scaled = self._current_pixmap.scaled(self.size(), Qt.KeepAspectRatio, Qt.SmoothTransformation)
            super().setPixmap(scaled)

    def resizeEvent(self, event):
        self.update_display()
        super().resizeEvent(event)


# ============================================================
# Background Render Workers
# ============================================================

class RenderSignals(QObject):
    thumb_ready = Signal(object)
    preview_ready = Signal(object)
    error = Signal(str)


class RenderTask(QRunnable):
    def __init__(self, signals: RenderSignals, payload: Dict[str, Any]):
        super().__init__()
        self.signals = signals
        self.payload = payload

    @staticmethod
    def _normalize_pixmap_rgb(pix: fitz.Pixmap) -> fitz.Pixmap:
        try:
            if pix.n != 3:
                return fitz.Pixmap(fitz.csRGB, pix)
        except Exception:
            pass
        return pix

    def run(self):
        try:
            kind = self.payload["kind"]
            key = self.payload["key"]
            render_path = self.payload["render_path"]
            page_idx = int(self.payload.get("page_idx", 0))
            target_px = int(self.payload.get("target_px", 160))
            token = self.payload.get("token", None)

            if not os.path.exists(render_path):
                raise FileNotFoundError(render_path)

            doc = fitz.open(render_path)
            try:
                if page_idx < 0 or page_idx >= len(doc):
                    raise IndexError(f"Invalid page index: {page_idx}")

                page = doc[page_idx]
                rect = page.rect
                if rect.width <= 0:
                    raise ValueError("Invalid page width")

                # target_px는 '가로폭' 기준으로 스케일 결정
                scale = max(0.05, target_px / float(rect.width))
                matrix = fitz.Matrix(scale, scale).prerotate(page.rotation)

                pix = page.get_pixmap(matrix=matrix, alpha=False)
                pix = self._normalize_pixmap_rgb(pix)

                out = {
                    "kind": kind,
                    "key": key,
                    "token": token,
                    "w": pix.width,
                    "h": pix.height,
                    "stride": pix.stride,
                    "n": pix.n,
                    "samples": bytes(pix.samples),
                }

                if kind == "thumb":
                    self.signals.thumb_ready.emit(out)
                else:
                    self.signals.preview_ready.emit(out)

            finally:
                doc.close()

        except Exception as e:
            self.signals.error.emit(str(e))


# ============================================================
# List Widget
# ============================================================

class MultiSelectListWidget(QListWidget):
    def __init__(self, parent=None, is_work: bool = False):
        super().__init__(parent)
        self.main_win = parent
        self.is_work = is_work

        self.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.setMouseTracking(True)
        self.setUniformItemSizes(False)

        self.setContextMenuPolicy(Qt.CustomContextMenu)
        self.customContextMenuRequested.connect(self.show_context_menu)

        self.setMinimumWidth(0)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

        # 아이템 간격 최소
        self.setStyleSheet(
            "QListWidget::item { padding: 0px; margin: 0px; }"
            "QListWidget::item:selected { background: #d0ebff; }"
        )

        if self.is_work:
            self.setDragEnabled(True)
            self.setAcceptDrops(True)
            self.viewport().setAcceptDrops(True)
            self.setDropIndicatorShown(True)
            self.setDragDropMode(QAbstractItemView.InternalMove)
            self.setDefaultDropAction(Qt.MoveAction)
        else:
            self.setDragEnabled(False)
            self.setAcceptDrops(True)
            self.viewport().setAcceptDrops(True)
            self.setDropIndicatorShown(True)
            self.setDragDropMode(QAbstractItemView.DropOnly)
            self.setDefaultDropAction(Qt.CopyAction)

    def mouseMoveEvent(self, event):
        item = self.itemAt(event.position().toPoint())
        if item:
            self.main_win.handle_item_preview(item)
        super().mouseMoveEvent(event)

    def wheelEvent(self, event):
        if self.count() == 0:
            super().wheelEvent(event)
            return
        current_row = self.currentRow()
        new_row = max(0, current_row - 1) if event.angleDelta().y() > 0 else min(self.count() - 1, current_row + 1)
        self.setCurrentRow(new_row)
        item = self.item(new_row)
        if item:
            self.main_win.handle_item_preview(item)
        event.accept()

    def dragEnterEvent(self, event):
        if self.is_work:
            if event.source() is self:
                event.acceptProposedAction()
            else:
                event.ignore()
            return

        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        if self.is_work:
            if event.source() is self:
                event.acceptProposedAction()
            else:
                event.ignore()
            return

        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        if self.is_work:
            if event.source() is self:
                super().dropEvent(event)
                self.main_win.on_worklist_reordered()
            else:
                event.ignore()
            return

        if event.mimeData().hasUrls():
            files = [u.toLocalFile() for u in event.mimeData().urls()]
            if files:
                self.main_win.load_files(files, clear_before=False, source="원본 리스트 드롭(파일 추가)")
            event.acceptProposedAction()
        else:
            event.ignore()

    def show_context_menu(self, pos):
        item = self.itemAt(pos)
        if not item:
            return
        if not item.isSelected():
            self.setCurrentItem(item)
        menu = QMenu()
        if not self.is_work:
            a1 = QAction("➕ 선택 항목 모두 추가", self)
            a1.triggered.connect(self.main_win.add_selected_to_work)
            menu.addAction(a1)
        else:
            a1 = QAction("❌ 선택 항목 모두 삭제", self)
            a1.triggered.connect(self.main_win.remove_work_items)
            menu.addAction(a1)
        menu.exec(QCursor.pos())


# ============================================================
# Main Window
# ============================================================

class IntegratedPDFMaster(QMainWindow):
    SUPPORTED_EXTS = {"pdf", "txt", "jpg", "jpeg", "png", "bmp"}

    PREVIEW_CACHE_LIMIT = 200
    PREVIEW_DEBOUNCE_MS = 80

    THUMB_QUEUE_LIMIT_PER_LOAD = 450
    THUMB_CACHE_LIMIT = 2500

    TEXT_MODE_PLAIN = "plain"
    TEXT_MODE_LAYOUT = "layout"

    CONVERT_MODE_AUTO = "auto"
    CONVERT_MODE_IMAGE = "image"
    CONVERT_MODE_MIXED = "mixed"

    MERGE_SIZE_KEEP = "keep"
    MERGE_SIZE_FIRST = "first"
    MERGE_SIZE_A4 = "a4"

    MERGE_ALIGN_CENTER = "center"
    MERGE_ALIGN_TOP = "top"

    MERGE_ROTATE_KEEP = "keep"
    MERGE_ROTATE_PORTRAIT = "portrait"
    MERGE_ROTATE_LANDSCAPE = "landscape"

    PDF_SAVE_DEFAULT = "default"
    PDF_SAVE_SMALL = "small"
    PDF_SAVE_PRINT = "print"

    def __init__(self):
        super().__init__()
        self.setWindowTitle("통합 PDF 마스터 Pro (ALL-IN-ONE)")
        self.setAcceptDrops(False)

        self.settings = QSettings("MySoft", "PDFMaster_AllInOne_Clean_Final")

        self.docs: List[SourceDoc] = []
        self.last_path = self.settings.value("last_path", os.path.expanduser("~"))
        self.current_dpi = int(self.settings.value("dpi", 150))
        self.current_font_size = int(self.settings.value("font_size", 11))
        self.text_import_font_size = int(self.settings.value("text_import_font_size", 11))
        self.text_mode = self.settings.value("text_mode", self.TEXT_MODE_PLAIN)
        self.convert_mode = self.settings.value("convert_mode", self.CONVERT_MODE_AUTO)

        self.font_path = self._find_korean_font_path()

        self.merge_size_mode = self.settings.value("merge_size_mode", self.MERGE_SIZE_KEEP)
        self.merge_align_mode = self.settings.value("merge_align_mode", self.MERGE_ALIGN_CENTER)
        self.merge_rotate_mode = self.settings.value("merge_rotate_mode", self.MERGE_ROTATE_KEEP)
        self.pdf_save_mode = self.settings.value("pdf_save_mode", self.PDF_SAVE_DEFAULT)

        # ✅ PDF 렌더 캐시(중복 생성 방지)
        self._pdf_render_cache: Dict[str, Dict[str, Any]] = {}

        # ✅ preview cache: (pixmap, rendered_width)
        self._preview_cache: "OrderedDict[Tuple[int, int], Tuple[QPixmap, int]]" = OrderedDict()
        self._thumb_cache: "OrderedDict[Tuple[int, int], QPixmap]" = OrderedDict()
        self._orig_item_map: Dict[Tuple[int, int], QListWidgetItem] = {}

        self._thumb_inflight: Set[Tuple[int, int]] = set()
        self._preview_inflight: Set[Tuple[int, int]] = set()

        self._pending_preview_key: Optional[Tuple[int, int]] = None
        self._current_preview_key: Optional[Tuple[int, int]] = None

        self._preview_timer = QTimer(self)
        self._preview_timer.setSingleShot(True)
        self._preview_timer.timeout.connect(self._render_pending_preview)

        self._work_key_set: Set[Tuple[int, int]] = set()

        self.thread_pool = QThreadPool.globalInstance()
        try:
            cpu = os.cpu_count() or 4
            self.thread_pool.setMaxThreadCount(max(2, min(8, cpu)))
        except Exception:
            pass

        self.render_signals = RenderSignals()
        self.render_signals.thumb_ready.connect(self._on_thumb_ready)
        self.render_signals.preview_ready.connect(self._on_preview_ready)
        self.render_signals.error.connect(self._on_render_error)

        self._last_render_error: Optional[str] = None

        self._icon_px: int = 48
        self.THUMB_MAX_W: int = self._icon_px
        self._thumb_generation: int = 0

        self.init_ui()
        self.create_full_menu_bar()
        self.load_settings()

    # ------------------------------------------------------------
    # Utilities
    # ------------------------------------------------------------
    def _ts(self) -> str:
        return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

    def default_save_path(self, base_name: str, ext: str) -> str:
        return os.path.join(self.last_path, f"{base_name}_{self._ts()}.{ext}")

    def update_last_path(self, path: str):
        if not path:
            return
        self.last_path = path if os.path.isdir(path) else os.path.dirname(path)
        self.settings.setValue("last_path", self.last_path)

    def set_status(self, t: str):
        self.status_bar.setText(f" ▶ {t}")

    def _find_korean_font_path(self) -> Optional[str]:
        base_dir = resource_base_dir()
        candidates = [
            os.path.join(base_dir, "fonts", "NotoSansKR-Regular.ttf"),
            os.path.join(base_dir, "fonts", "NotoSansKR-Regular.otf"),
            os.path.join(base_dir, "fonts", "NanumGothic.ttf"),
            r"C:\Windows\Fonts\malgun.ttf",
            r"C:\Windows\Fonts\malgunbd.ttf",
            r"C:\Windows\Fonts\gulim.ttc",
            "/System/Library/Fonts/AppleGothic.ttf",
            "/System/Library/Fonts/Supplemental/AppleGothic.ttf",
            "/Library/Fonts/AppleGothic.ttf",
            "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
            "/usr/share/fonts/truetype/noto/NotoSansKR-Regular.otf",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        ]
        for p in candidates:
            if p and os.path.exists(p):
                return p
        return None

    # ------------------------------------------------------------
    # ✅ High-Res Preview helpers (DPR 기반)
    # ------------------------------------------------------------
    def _device_pixel_ratio(self) -> float:
        try:
            wh = self.windowHandle()
            if wh and wh.screen():
                return float(wh.screen().devicePixelRatio())
        except Exception:
            pass
        try:
            s = QApplication.primaryScreen()
            if s:
                return float(s.devicePixelRatio())
        except Exception:
            pass
        return 1.0

    def _desired_preview_target_px(self) -> int:
        # 프리뷰 표시 영역(라벨 폭) 기반
        w = max(400, int(self.preview_label.width()))
        dpr = self._device_pixel_ratio()

        # 모드별 품질 계수/상한
        if self.pdf_save_mode == self.PDF_SAVE_SMALL:
            quality = 1.10
            cap = 1800
            floor = 900
        elif self.pdf_save_mode == self.PDF_SAVE_PRINT:
            quality = 1.55
            cap = 4200
            floor = 1400
        else:
            quality = 1.35
            cap = 2800
            floor = 1200

        target = int(w * dpr * quality)
        return max(floor, min(cap, target))

    # ------------------------------------------------------------
    # ✅ PDF 캐시 중복 방지
    # ------------------------------------------------------------
    def _pdf_signature(self, pdf_path: str) -> str:
        ap = os.path.abspath(pdf_path)
        try:
            st = os.stat(ap)
            mtime = int(st.st_mtime)
            size = int(st.st_size)
        except Exception:
            mtime = 0
            size = 0
        return f"{ap}|{mtime}|{size}"

    def _cache_pdf_for_render(self, pdf_path: str) -> str:
        sig = self._pdf_signature(pdf_path)
        rec = self._pdf_render_cache.get(sig)
        if rec and os.path.exists(rec.get("path", "")):
            rec["ref"] = int(rec.get("ref", 0)) + 1
            rec["last"] = datetime.datetime.now().isoformat(timespec="seconds")
            return rec["path"]

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf", prefix="pdf_render_")
        tmp_path = tmp.name
        tmp.close()

        try:
            src = fitz.open(pdf_path)
            try:
                src.save(tmp_path, garbage=4, deflate=True)
            finally:
                src.close()
        except Exception:
            try:
                shutil.copy2(pdf_path, tmp_path)
            except Exception:
                try:
                    os.remove(tmp_path)
                except Exception:
                    pass
                return pdf_path

        self._pdf_render_cache[sig] = {
            "path": tmp_path,
            "ref": 1,
            "src": os.path.abspath(pdf_path),
            "last": datetime.datetime.now().isoformat(timespec="seconds"),
        }
        return tmp_path

    # ------------------------------------------------------------
    # ✅ TXT -> PDF (fpdf2 조판, uni=True 폰트 등록)
    # ------------------------------------------------------------
    def _read_text_file(self, path: str) -> str:
        for enc in ("utf-8-sig", "utf-8", "cp949", "euc-kr", "latin1"):
            try:
                with open(path, "r", encoding=enc) as f:
                    return f.read()
            except Exception:
                continue
        raise RuntimeError("지원 인코딩(UTF-8/CP949/EUC-KR)으로 읽을 수 없습니다.")

    def _txt_to_cached_pdf(self, txt_path: str) -> str:
        if FPDF is None:
            raise RuntimeError("TXT → PDF 변환을 위해 fpdf2가 필요합니다.\n\n pip install --upgrade fpdf2")

        content = self._read_text_file(txt_path)
        if not content.strip():
            pdf = FPDF(format="A4")
            pdf.add_page()
            out_bytes = pdf.output(dest="S")
            if isinstance(out_bytes, str):
                out_bytes = out_bytes.encode("latin-1", errors="ignore")
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf", prefix="txt_render_")
            tmp.write(out_bytes)
            tmp.close()
            return tmp.name

        body_pt = int(self.text_import_font_size)
        chapter_pt = max(body_pt + 4, 14)
        line_spacing = 1.8
        para_gap = 2
        indent_spaces = 4

        pdf = FPDF(format="A4")
        pdf.set_auto_page_break(True, margin=15)

        family = "Helvetica"
        if self.font_path and os.path.exists(self.font_path):
            try:
                family = "NovelFont"
                # ✅ uni=True: 한글 유니코드 안정
                pdf.add_font(family=family, style="", fname=self.font_path, uni=True)
                pdf.add_font(family=family, style="B", fname=self.font_path, uni=True)
            except Exception:
                family = "Helvetica"

        if contains_korean(content) and family == "Helvetica":
            raise RuntimeError(
                "한글 본문이 감지되었지만 PDF 폰트 등록에 실패했습니다.\n"
                "- fonts/NotoSansKR-Regular.ttf(권장)를 동봉하거나\n"
                "- 시스템에 한글 TTF 폰트(예: malgun.ttf)가 존재하는지 확인하세요.\n"
                f"- 현재 탐색된 폰트 경로: {self.font_path or '(없음)'}"
            )

        chapters = split_into_chapters(content, strict=True, min_body_lines=3)
        line_h = max(4.0, line_spacing * (body_pt * 0.45))

        for ch in chapters:
            pdf.add_page()
            pdf.set_font(family, style="B", size=chapter_pt)
            pdf.cell(0, 10, txt=ch["title"], align="C", ln=1)
            pdf.ln(2)

            pdf.set_font(family, style="", size=body_pt)
            for p in smart_make_paragraphs(ch["lines"]):
                if not p.strip():
                    pdf.ln(4)
                    continue
                stripped = p.lstrip()
                is_dlg = stripped.startswith(DIALOGUE_PREFIXES)
                txt = stripped if is_dlg else (" " * indent_spaces + stripped)
                pdf.multi_cell(0, line_h, txt=txt, align="J")
                if para_gap:
                    pdf.ln(para_gap * 0.5)

        out_bytes = pdf.output(dest="S")
        if isinstance(out_bytes, str):
            out_bytes = out_bytes.encode("latin-1", errors="ignore")

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf", prefix="txt_render_")
        tmp.write(out_bytes)
        tmp.close()
        return tmp.name

    # ------------------------------------------------------------
    # UI / list sizing
    # ------------------------------------------------------------
    def _update_list_icon_size(self):
        ui_font = QApplication.font()
        fm = QFontMetrics(ui_font)
        text_h = fm.height()
        THUMB_SCALE = 1.55
        new_icon_px = max(24, int(text_h * THUMB_SCALE))
        changed = (new_icon_px != self._icon_px)
        if changed:
            self._thumb_generation += 1
        self._icon_px = new_icon_px
        self.THUMB_MAX_W = new_icon_px
        size = QSize(self._icon_px, self._icon_px)
        row_h = max(self._icon_px, fm.height()) + 2
        for lw in (self.orig_list, self.work_list):
            lw.setFont(ui_font)
            lw.setIconSize(size)
            lw.setSpacing(0)
            for i in range(lw.count()):
                it = lw.item(i)
                if it:
                    it.setFont(ui_font)
                    it.setSizeHint(QSize(0, row_h))
        if changed:
            self._thumb_cache.clear()
            self._thumb_inflight.clear()

    # ------------------------------------------------------------
    # UI
    # ------------------------------------------------------------
    def init_ui(self):
        self.central_widget = QWidget()
        self.setCentralWidget(self.central_widget)

        self.main_layout = QVBoxLayout(self.central_widget)
        self.main_layout.setContentsMargins(0, 0, 0, 0)
        self.main_layout.setSpacing(0)

        self.tool_panel = QFrame()
        self.tool_panel.setFixedHeight(70)
        self.tool_panel.setStyleSheet("QFrame { background: #f8f9fa; border-bottom: 1px solid #dee2e6; }")
        self.tool_layout = QHBoxLayout(self.tool_panel)
        self.create_tool_buttons()
        self.main_layout.addWidget(self.tool_panel)

        self.content_splitter = QSplitter(Qt.Horizontal)
        self.content_splitter.setHandleWidth(6)
        self.content_splitter.setChildrenCollapsible(True)

        self.orig_panel = self._make_list_panel("📂 원본 페이지", is_work=False)
        self.content_splitter.addWidget(self.orig_panel)

        self.work_panel = self._make_list_panel("📑 작업 리스트", is_work=True)
        self.content_splitter.addWidget(self.work_panel)

        w3 = QWidget()
        l3 = QVBoxLayout(w3)
        l3.setContentsMargins(9, 9, 9, 9)
        title = QLabel("<b>🔍 실시간 미리보기</b>")
        self.preview_label = ResponsiveLabel()
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setWidget(self.preview_label)
        l3.addWidget(title)
        l3.addWidget(scroll)
        self.content_splitter.addWidget(w3)

        self.content_splitter.setStretchFactor(0, 1)
        self.content_splitter.setStretchFactor(1, 1)
        self.content_splitter.setStretchFactor(2, 2)

        self.main_layout.addWidget(self.content_splitter)

        self.status_container = QFrame()
        self.status_container.setFixedHeight(30)
        self.status_container.setStyleSheet("background: #343a40;")

        self.status_layout = QHBoxLayout(self.status_container)
        self.status_layout.setContentsMargins(10, 0, 10, 0)

        self.status_bar = QLabel(" 시스템 준비 완료")
        self.status_bar.setStyleSheet("color: #00ff88;")

        self.progress_bar = QProgressBar()
        self.progress_bar.setFixedWidth(260)
        self.progress_bar.setFixedHeight(15)
        self.progress_bar.setVisible(False)
        self.progress_bar.setTextVisible(True)
        self.progress_bar.setStyleSheet(
            "QProgressBar { background: #495057; border-radius: 5px; color: white; text-align: center; } "
            "QProgressBar::chunk { background: #00ff88; }"
        )

        self.status_layout.addWidget(self.status_bar)
        self.status_layout.addStretch()
        self.status_layout.addWidget(self.progress_bar)
        self.main_layout.addWidget(self.status_container)

    def _make_list_panel(self, title: str, is_work: bool) -> QWidget:
        w = QWidget()
        l = QVBoxLayout(w)
        l.setContentsMargins(9, 9, 9, 9)
        l.setSpacing(6)

        title_label = QLabel(f"<b>{title}</b>")
        l.addWidget(title_label)

        search = QLineEdit()
        search.setPlaceholderText("검색(필터): 파일명 / 페이지 ...")
        search.setClearButtonEnabled(True)
        l.addWidget(search)

        list_widget = MultiSelectListWidget(self, is_work=is_work)
        list_widget.itemClicked.connect(self.handle_item_preview)

        if is_work:
            self.work_list = list_widget
            self.work_search = search
            self.work_list.itemDoubleClicked.connect(self.remove_work_items)
            search.textChanged.connect(lambda txt: self._apply_filter(self.work_list, txt))

            btn = QPushButton("리스트에서 삭제")
            btn.setStyleSheet("background: #ff4757; color: white; font-weight: bold;")
            btn.clicked.connect(self.remove_work_items)

            l.addWidget(list_widget)
            l.addWidget(btn)

            tip = QLabel("<span style='color:#666'>TIP: 작업 리스트는 드래그로 순서를 바꿀 수 있습니다.</span>")
            tip.setWordWrap(True)
            l.addWidget(tip)
        else:
            self.orig_list = list_widget
            self.orig_search = search
            self.orig_list.itemDoubleClicked.connect(self.add_selected_to_work)
            search.textChanged.connect(lambda txt: self._apply_filter(self.orig_list, txt))
            l.addWidget(list_widget)

        return w

    def _apply_filter(self, lw: QListWidget, text: str):
        t = (text or "").strip().lower()
        for i in range(lw.count()):
            it = lw.item(i)
            it.setHidden(False if not t else (t not in it.text().lower()))

    def create_tool_buttons(self):
        btns = [
            ("열기", "📂", lambda: self.open_files(True)),
            ("추가", "➕", lambda: self.open_files(False)),
            ("초기화", "🧹", self.clear_all),
            ("|", None, None),
            ("병합", "🔗", self.action_merge),
            ("추출", "✂", self.action_extract),
            ("|", None, None),
            ("Word", "📝", self.convert_to_word),
            ("PPT", "📊", self.convert_to_ppt),
            ("|", None, None),
            ("Excel", "📑", self.convert_to_excel),
            ("Text", "📄", self.convert_to_text),
            ("이미지", "🖼", self.convert_to_images),
        ]
        for t, i, f in btns:
            if t == "|":
                line = QFrame()
                line.setFrameShape(QFrame.VLine)
                line.setStyleSheet("color: #ced4da; margin: 10px 5px;")
                self.tool_layout.addWidget(line)
                continue
            btn = QPushButton(f"{i}\n{t}")
            btn.setFixedSize(68, 48)
            btn.setStyleSheet(
                "QPushButton { background: transparent; border-radius: 5px; font-weight: bold; } "
                "QPushButton:hover { background: #e9ecef; }"
            )
            btn.clicked.connect(f)
            self.tool_layout.addWidget(btn)
        self.tool_layout.addStretch()

    def create_full_menu_bar(self):
        mb = self.menuBar()

        file_m = mb.addMenu("파일(&F)")
        file_m.addAction("새로 열기").triggered.connect(lambda: self.open_files(True))
        file_m.addAction("파일 추가").triggered.connect(lambda: self.open_files(False))
        file_m.addAction("전체 초기화").triggered.connect(self.clear_all)
        file_m.addSeparator()
        file_m.addAction("종료").triggered.connect(self.close)

        pdf_m = mb.addMenu("PDF 생성(&P)")
        pdf_m.addAction("작업 리스트 병합 저장").triggered.connect(self.action_merge)
        pdf_m.addAction("작업 리스트 개별 추출").triggered.connect(self.action_extract)

        conv_m = mb.addMenu("변환(&C)")
        conv_m.addAction("Word").triggered.connect(self.convert_to_word)
        conv_m.addAction("PPT").triggered.connect(self.convert_to_ppt)
        conv_m.addAction("Excel").triggered.connect(self.convert_to_excel)
        conv_m.addAction("Text").triggered.connect(self.convert_to_text)
        conv_m.addAction("이미지").triggered.connect(self.convert_to_images)

        set_m = mb.addMenu("설정(&S)")

        # TXT 임포트 폰트 크기
        txt_f_menu = set_m.addMenu("Text 가져오기 폰트 크기")
        self.txt_font_group = QActionGroup(self)
        for s in [9, 10, 11, 12, 13, 14]:
            a = QAction(f"{s}pt", self, checkable=True)
            a.setChecked(s == self.text_import_font_size)
            a.triggered.connect(lambda chk=False, sz=s: self.change_text_import_font(sz))
            self.txt_font_group.addAction(a)
            txt_f_menu.addAction(a)

        # 이미지 저장 해상도
        res_menu = set_m.addMenu("이미지 저장 해상도")
        self.res_group = QActionGroup(self)
        for label, dpi in [("72 DPI", 72), ("150 DPI", 150), ("300 DPI", 300), ("600 DPI", 600)]:
            a = QAction(label, self, checkable=True)
            a.setChecked(dpi == self.current_dpi)
            a.triggered.connect(lambda chk=False, d=dpi: self.change_resolution(d))
            self.res_group.addAction(a)
            res_menu.addAction(a)

        # UI 글자 크기
        f_menu = set_m.addMenu("UI 글자 크기")
        self.ui_font_group = QActionGroup(self)
        for s in [10, 11, 12, 14, 16]:
            a = QAction(f"{s}pt", self, checkable=True)
            a.setChecked(s == self.current_font_size)
            a.triggered.connect(lambda chk=False, sz=s: self.change_ui_font(sz))
            self.ui_font_group.addAction(a)
            f_menu.addAction(a)

        # 텍스트 추출 모드
        text_mode_menu = set_m.addMenu("텍스트 추출 모드")
        self.text_mode_group = QActionGroup(self)
        for label, mode in [("일반", self.TEXT_MODE_PLAIN), ("레이아웃 유지(정렬)", self.TEXT_MODE_LAYOUT)]:
            a = QAction(label, self, checkable=True)
            a.setChecked(mode == self.text_mode)
            a.triggered.connect(lambda chk=False, m=mode: self.change_text_mode(m))
            self.text_mode_group.addAction(a)
            text_mode_menu.addAction(a)

        # 변환 모드
        convert_mode_menu = set_m.addMenu("변환 모드(Word/PPT/Excel)")
        self.convert_mode_group = QActionGroup(self)
        for label, mode in [
            ("자동(텍스트 우선)", self.CONVERT_MODE_AUTO),
            ("이미지 우선", self.CONVERT_MODE_IMAGE),
            ("혼합(이미지 + 텍스트)", self.CONVERT_MODE_MIXED),
        ]:
            a = QAction(label, self, checkable=True)
            a.setChecked(mode == self.convert_mode)
            a.triggered.connect(lambda chk=False, m=mode: self.change_convert_mode(m))
            self.convert_mode_group.addAction(a)
            convert_mode_menu.addAction(a)

        # PDF 저장 모드
        pdf_mode_menu = set_m.addMenu("PDF 저장 모드")
        self.pdf_save_mode_group = QActionGroup(self)
        for label, mode in [
            ("일반 (권장)", self.PDF_SAVE_DEFAULT),
            ("저용량 (화면용)", self.PDF_SAVE_SMALL),
            ("고품질 (인쇄용)", self.PDF_SAVE_PRINT),
        ]:
            a = QAction(label, self, checkable=True)
            a.setChecked(mode == self.pdf_save_mode)
            a.triggered.connect(lambda chk=False, m=mode: self.change_pdf_save_mode(m))
            self.pdf_save_mode_group.addAction(a)
            pdf_mode_menu.addAction(a)

        # 병합 옵션
        merge_menu = set_m.addMenu("PDF 병합 옵션")

        size_menu = merge_menu.addMenu("페이지 크기")
        self.merge_size_group = QActionGroup(self)
        for label, mode in [
            ("원본 유지", self.MERGE_SIZE_KEEP),
            ("첫 페이지 기준", self.MERGE_SIZE_FIRST),
            ("A4로 통일", self.MERGE_SIZE_A4),
        ]:
            a = QAction(label, self, checkable=True)
            a.setChecked(self.merge_size_mode == mode)
            a.triggered.connect(lambda chk=False, m=mode: setattr(self, "merge_size_mode", m))
            self.merge_size_group.addAction(a)
            size_menu.addAction(a)

        align_menu = merge_menu.addMenu("페이지 배치")
        self.merge_align_group = QActionGroup(self)
        for label, mode in [("중앙 정렬", self.MERGE_ALIGN_CENTER), ("상단 정렬", self.MERGE_ALIGN_TOP)]:
            a = QAction(label, self, checkable=True)
            a.setChecked(self.merge_align_mode == mode)
            a.triggered.connect(lambda chk=False, m=mode: setattr(self, "merge_align_mode", m))
            self.merge_align_group.addAction(a)
            align_menu.addAction(a)

        rotate_menu = merge_menu.addMenu("자동 회전")
        self.merge_rotate_group = QActionGroup(self)
        for label, mode in [
            ("원본 유지", self.MERGE_ROTATE_KEEP),
            ("세로로 통일", self.MERGE_ROTATE_PORTRAIT),
            ("가로로 통일", self.MERGE_ROTATE_LANDSCAPE),
        ]:
            a = QAction(label, self, checkable=True)
            a.setChecked(self.merge_rotate_mode == mode)
            a.triggered.connect(lambda chk=False, m=mode: setattr(self, "merge_rotate_mode", m))
            self.merge_rotate_group.addAction(a)
            rotate_menu.addAction(a)

    # ------------------------------------------------------------
    # Settings changes
    # ------------------------------------------------------------
    def change_text_import_font(self, size: int):
        self.text_import_font_size = size
        self.set_status(f"Text 임포트 폰트: {size}pt")

    def change_ui_font(self, size: int):
        self.current_font_size = size
        families = set(QFontDatabase.families())
        if "Malgun Gothic" in families:
            f = QFont("Malgun Gothic", size)
        else:
            f = QApplication.font()
            f.setPointSize(size)
        QApplication.setFont(f)
        self._update_list_icon_size()
        self.set_status(f"UI 글자 크기: {size}pt")

    def change_resolution(self, dpi: int):
        self.current_dpi = dpi
        self.set_status(f"해상도: {dpi} DPI")

    def change_text_mode(self, mode: str):
        self.text_mode = mode
        self.set_status(f"텍스트 추출 모드: {'레이아웃 유지' if mode == self.TEXT_MODE_LAYOUT else '일반'}")

    def change_convert_mode(self, mode: str):
        self.convert_mode = mode
        label = {
            self.CONVERT_MODE_AUTO: "자동(텍스트 우선)",
            self.CONVERT_MODE_IMAGE: "이미지 우선",
            self.CONVERT_MODE_MIXED: "혼합(이미지+텍스트)",
        }.get(mode, mode)
        self.set_status(f"변환 모드: {label}")

    def change_pdf_save_mode(self, mode: str):
        self.pdf_save_mode = mode
        label = {
            self.PDF_SAVE_DEFAULT: "일반",
            self.PDF_SAVE_SMALL: "저용량",
            self.PDF_SAVE_PRINT: "고품질",
        }.get(mode, mode)
        try:
            self.settings.setValue("pdf_save_mode", self.pdf_save_mode)
        except Exception:
            pass
        self.set_status(f"PDF 저장 모드: {label} / 미리보기: {'인쇄용' if mode==self.PDF_SAVE_PRINT else '화면용'}")

        # ✅ 모드 전환 시: 프리뷰 캐시가 너무 저해상도면 자동 재렌더 되지만,
        # 품질 차이를 바로 반영하고 싶다면 현재 프리뷰만 즉시 재요청
        if self._current_preview_key:
            self._pending_preview_key = self._current_preview_key
            self._preview_timer.start(10)

    # ------------------------------------------------------------
    # PDF save helper (모드별)
    # ------------------------------------------------------------
    def _save_pdf(self, doc: fitz.Document, save_path: str):
        if self.pdf_save_mode == self.PDF_SAVE_SMALL:
            self.save_pdf_low_quality(doc, save_path, dpi=120, jpg_quality=70)
        elif self.pdf_save_mode == self.PDF_SAVE_PRINT:
            doc.save(save_path, garbage=1, deflate=False)
        else:
            doc.save(save_path, garbage=4, deflate=True)

    def save_pdf_low_quality(self, src_doc: fitz.Document, save_path: str, dpi: int = 120, jpg_quality: int = 70):
        out = fitz.open()
        zoom = dpi / 72.0
        for i in range(src_doc.page_count):
            page = src_doc[i]
            rect = page.rect
            matrix = fitz.Matrix(zoom, zoom).prerotate(page.rotation)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            if pix.n != 3:
                pix = fitz.Pixmap(fitz.csRGB, pix)
            new_page = out.new_page(width=rect.width, height=rect.height)
            dst = new_page.rect
            img_w, img_h = pix.width, pix.height
            dst_w, dst_h = dst.width, dst.height
            img_ratio = img_w / img_h if img_h else 1.0
            dst_ratio = dst_w / dst_h if dst_h else 1.0
            if img_ratio > dst_ratio:
                w = dst_w
                h = dst_w / img_ratio
            else:
                h = dst_h
                w = dst_h * img_ratio
            x0 = (dst_w - w) / 2
            y0 = (dst_h - h) / 2
            place = fitz.Rect(x0, y0, x0 + w, y0 + h)
            new_page.insert_image(place, stream=pix.tobytes("jpg", jpg_quality=jpg_quality))
        out.save(save_path, garbage=4, deflate=True)
        out.close()

    # ------------------------------------------------------------
    # File open/add
    # ------------------------------------------------------------
    def open_files(self, clear: bool):
        paths, _ = QFileDialog.getOpenFileNames(
            self,
            "파일 열기",
            self.last_path,
            "지원 파일 (*.pdf *.txt *.jpg *.jpeg *.png *.bmp)"
        )
        if paths:
            self.load_files(paths, clear_before=clear, source="파일 대화상자")

    def _write_temp_pdf(self, pdf_bytes: bytes, prefix: str) -> str:
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf", prefix=prefix)
        tmp.write(pdf_bytes)
        tmp.close()
        return tmp.name

    def _image_to_pdf_bytes(self, path: str) -> bytes:
        ext = os.path.splitext(path)[1].lower().lstrip(".")
        with open(path, "rb") as bf:
            raw = bf.read()
        img_doc = fitz.open(stream=raw, filetype=ext)
        try:
            pdf_bytes = img_doc.convert_to_pdf()
        finally:
            img_doc.close()
        return pdf_bytes

    def open_source_document(self, path: str) -> Optional[SourceDoc]:
        base = os.path.basename(path)
        ext = os.path.splitext(path)[1].lower().lstrip(".")

        doc: Optional[fitz.Document] = None
        source_type: Optional[str] = None
        render_path: Optional[str] = None
        temp_paths: List[str] = []

        if ext == "pdf":
            source_type = "pdf"
            render_path = self._cache_pdf_for_render(path)
            doc = fitz.open(render_path)

        elif ext == "txt":
            source_type = "txt"
            render_path = self._txt_to_cached_pdf(path)
            temp_paths.append(render_path)
            doc = fitz.open(render_path)

        elif ext in ["jpg", "jpeg", "png", "bmp"]:
            source_type = "image"
            pdf_bytes = self._image_to_pdf_bytes(path)
            render_path = self._write_temp_pdf(pdf_bytes, prefix="img_render_")
            temp_paths.append(render_path)
            doc = fitz.open(render_path)

        if doc is None or source_type is None or render_path is None:
            return None

        return SourceDoc(doc=doc, name=base, path=path, source_type=source_type,
                         render_path=render_path, temp_paths=temp_paths)

    def load_files(self, files: List[str], clear_before: bool = False, source: str = ""):
        files = [f for f in files if os.path.isfile(f) and f.split(".")[-1].lower() in self.SUPPORTED_EXTS]
        if not files:
            self.set_status("지원하지 않는 파일 형식이 포함되어 있습니다.")
            return

        if clear_before:
            self.clear_all()
            self.set_status(f"새로 열기 모드 ({source})")
        else:
            self.set_status(f"파일 추가 모드 ({source})")

        self.update_last_path(files[0])

        total = len(files)
        self.progress_bar.setValue(0)
        self.progress_bar.setMaximum(total)
        self.progress_bar.setFormat("파일 로드 중...")
        self.progress_bar.setVisible(True)

        added_files = 0
        added_pages = 0
        thumb_queue_count = 0

        try:
            for idx, path in enumerate(files):
                base = os.path.basename(path)
                try:
                    self.set_status(f"파일 처리 중: {base}")
                    QApplication.processEvents()

                    sd = self.open_source_document(path)
                    if sd is None:
                        continue

                    d_id = len(self.docs)
                    self.docs.append(sd)

                    pages = len(sd.doc)
                    added_files += 1
                    added_pages += pages

                    for p_idx in range(pages):
                        it = QListWidgetItem(f"[{sd.name}] P.{p_idx + 1}")
                        it.setData(Qt.UserRole, (d_id, p_idx))
                        self.orig_list.addItem(it)
                        self._orig_item_map[(d_id, p_idx)] = it

                        if thumb_queue_count < self.THUMB_QUEUE_LIMIT_PER_LOAD:
                            self._enqueue_thumbnail_task(d_id, p_idx)
                            thumb_queue_count += 1

                except Exception as e:
                    QMessageBox.warning(self, "오류", f"{base}\n\n{e}")

                self.progress_bar.setValue(idx + 1)
                QApplication.processEvents()

        finally:
            self.progress_bar.setVisible(False)

        self._update_list_icon_size()
        self.set_status(f"로드 완료: 파일 {added_files}개 / 페이지 {added_pages}개")

    # ------------------------------------------------------------
    # Background thumbs / preview
    # ------------------------------------------------------------
    def _thumb_cache_put(self, key: Tuple[int, int], pm: QPixmap):
        self._thumb_cache[key] = pm
        self._thumb_cache.move_to_end(key)
        if len(self._thumb_cache) > self.THUMB_CACHE_LIMIT:
            self._thumb_cache.popitem(last=False)

    def _enqueue_thumbnail_task(self, d_id: int, p_idx: int):
        key = (d_id, p_idx)
        if key in self._thumb_inflight:
            return
        self._thumb_inflight.add(key)
        meta = self.docs[d_id]
        payload = {
            "kind": "thumb",
            "key": key,
            "render_path": meta.render_path,
            "page_idx": p_idx,
            "target_px": int(self.THUMB_MAX_W),
            "token": self._thumb_generation,
        }
        self.thread_pool.start(RenderTask(self.render_signals, payload))

    def _enqueue_preview_task(self, d_id: int, p_idx: int):
        key = (d_id, p_idx)
        if key in self._preview_inflight:
            return
        self._preview_inflight.add(key)
        meta = self.docs[d_id]
        payload = {
            "kind": "preview",
            "key": key,
            "render_path": meta.render_path,
            "page_idx": p_idx,
            # ✅ 고해상도 목표 폭(DPR 기반)
            "target_px": int(self._desired_preview_target_px()),
        }
        self.thread_pool.start(RenderTask(self.render_signals, payload))

    def _qimage_from_render_out(self, out: Dict[str, Any]):
        from PySide6.QtGui import QImage
        w, h, stride, n = out["w"], out["h"], out["stride"], out["n"]
        samples = out["samples"]
        if n == 1:
            fmt = QImage.Format_Grayscale8
        elif n == 4:
            fmt = QImage.Format_RGBA8888
        else:
            fmt = QImage.Format_RGB888
        return QImage(samples, w, h, stride, fmt).copy()

    @Slot(object)
    def _on_thumb_ready(self, out: Dict[str, Any]):
        key = out["key"]
        self._thumb_inflight.discard(key)
        token = out.get("token", None)
        if token is not None and token != self._thumb_generation:
            return
        pm = QPixmap.fromImage(self._qimage_from_render_out(out))
        self._thumb_cache_put(key, pm)
        it = self._orig_item_map.get(key)
        if it:
            it.setIcon(QIcon(pm))
            it.setSizeHint(QSize(0, self._icon_px + 2))
        for i in range(self.work_list.count()):
            wit = self.work_list.item(i)
            if wit and wit.data(Qt.UserRole) == key:
                wit.setIcon(QIcon(pm))
                wit.setSizeHint(QSize(0, self._icon_px + 2))

    @Slot(object)
    def _on_preview_ready(self, out: Dict[str, Any]):
        key = out["key"]
        self._preview_inflight.discard(key)
        pm = QPixmap.fromImage(self._qimage_from_render_out(out))

        rendered_w = int(out.get("w", pm.width()))
        self._cache_put_preview(key, pm, rendered_w)

        if self._current_preview_key == key:
            self.preview_label.set_pixmap_fixed(pm)

    @Slot(str)
    def _on_render_error(self, msg: str):
        if msg and msg != self._last_render_error:
            self._last_render_error = msg
            self.set_status(f"렌더 오류(마지막): {msg}")

    # ------------------------------------------------------------
    # Preview
    # ------------------------------------------------------------
    def handle_item_preview(self, item: QListWidgetItem):
        if not item:
            return
        key = item.data(Qt.UserRole)
        self._pending_preview_key = key
        try:
            d_id, p_idx = key
            if key not in self._thumb_cache and key not in self._thumb_inflight:
                self._enqueue_thumbnail_task(d_id, p_idx)
        except Exception:
            pass
        self._preview_timer.start(self.PREVIEW_DEBOUNCE_MS)

    def _cache_put_preview(self, key: Tuple[int, int], pm: QPixmap, rendered_w: int):
        self._preview_cache[key] = (pm, int(rendered_w))
        self._preview_cache.move_to_end(key)
        if len(self._preview_cache) > self.PREVIEW_CACHE_LIMIT:
            self._preview_cache.popitem(last=False)

    def _render_pending_preview(self):
        key = self._pending_preview_key
        if not key:
            return
        self._current_preview_key = key

        need = int(self._desired_preview_target_px())

        if key in self._preview_cache:
            pm, rendered_w = self._preview_cache[key]
            self._preview_cache.move_to_end(key)
            self.preview_label.set_pixmap_fixed(pm)

            # ✅ 더 큰 해상도가 필요하면 자동 재렌더
            if rendered_w < int(need * 0.92):
                d_id, p_idx = key
                self._enqueue_preview_task(d_id, p_idx)
            return

        if key in self._thumb_cache:
            self.preview_label.set_pixmap_fixed(self._thumb_cache[key])

        d_id, p_idx = key
        self._enqueue_preview_task(d_id, p_idx)

    # ------------------------------------------------------------
    # Work list
    # ------------------------------------------------------------
    def on_worklist_reordered(self):
        self.set_status("작업 리스트 순서 변경됨")

    def add_selected_to_work(self):
        selected = self.orig_list.selectedItems()
        added = 0
        for it in selected:
            key = it.data(Qt.UserRole)
            if key in self._work_key_set:
                continue
            ni = QListWidgetItem(it.text())
            ni.setData(Qt.UserRole, key)
            if key in self._thumb_cache:
                ni.setIcon(QIcon(self._thumb_cache[key]))
            else:
                try:
                    d_id, p_idx = key
                    self._enqueue_thumbnail_task(d_id, p_idx)
                except Exception:
                    pass
            self.work_list.addItem(ni)
            self._work_key_set.add(key)
            added += 1
        if selected:
            self.set_status(f"{added}개 항목 추가됨 (중복 제외)")
        self._update_list_icon_size()

    def remove_work_items(self):
        selected = self.work_list.selectedItems()
        removed = 0
        for it in selected:
            key = it.data(Qt.UserRole)
            self._work_key_set.discard(key)
            self.work_list.takeItem(self.work_list.row(it))
            removed += 1
        if selected:
            self.set_status(f"{removed}개 항목 삭제됨")

    # ------------------------------------------------------------
    # Reset
    # ------------------------------------------------------------
    def clear_all(self):
        for sd in self.docs:
            sd.close()
            sd.cleanup()
        self.docs = []
        self._preview_cache.clear()
        self._thumb_cache.clear()
        self._work_key_set.clear()
        self._orig_item_map.clear()
        self._thumb_inflight.clear()
        self._preview_inflight.clear()
        self.orig_list.clear()
        self.work_list.clear()
        self.preview_label.clear()
        self.set_status("초기화 완료")

    # ------------------------------------------------------------
    # Text extraction / rendering
    # ------------------------------------------------------------
    def extract_text(self, page: fitz.Page) -> str:
        try:
            if self.text_mode == self.TEXT_MODE_LAYOUT:
                return (page.get_text("text", sort=True) or "").strip()
            return (page.get_text() or "").strip()
        except Exception:
            return ""

    def render_page_image_bytes(self, page: fitz.Page, fmt: str = "png") -> io.BytesIO:
        fmt = (fmt or "png").lower()
        if fmt == "jpeg":
            fmt = "jpg"
        zoom = self.current_dpi / 72.0
        matrix = fitz.Matrix(zoom, zoom).prerotate(page.rotation)
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        try:
            if pix.n != 3:
                pix = fitz.Pixmap(fitz.csRGB, pix)
        except Exception:
            pass
        bio = io.BytesIO(pix.tobytes(fmt))
        bio.seek(0)
        return bio

    # ------------------------------------------------------------
    # Merge / Extract
    # ------------------------------------------------------------
    def action_merge(self):
        if self.work_list.count() == 0:
            return
        default_file = self.default_save_path("병합완료_결과물", "pdf")
        save_path, _ = QFileDialog.getSaveFileName(self, "PDF 병합 저장", default_file, "PDF (*.pdf)")
        if not save_path:
            return
        self.update_last_path(save_path)

        out = fitz.open()
        target_w = target_h = None
        if self.merge_size_mode == self.MERGE_SIZE_FIRST:
            d0, i0 = self.work_list.item(0).data(Qt.UserRole)
            r0 = self.docs[d0].doc[i0].rect
            target_w, target_h = r0.width, r0.height
        elif self.merge_size_mode == self.MERGE_SIZE_A4:
            target_w, target_h = 595, 842

        CONTENT_MARGIN_LR = 55
        CONTENT_MARGIN_TOP = 65
        CONTENT_MARGIN_BOTTOM = 65

        for k in range(self.work_list.count()):
            d, idx = self.work_list.item(k).data(Qt.UserRole)
            meta = self.docs[d]
            src_doc = meta.doc
            src_page = src_doc[idx]
            rect = src_page.rect
            sw, sh = rect.width, rect.height

            rotation = 0
            if self.merge_rotate_mode == self.MERGE_ROTATE_PORTRAIT and sw > sh:
                rotation = 90
                sw, sh = sh, sw
            elif self.merge_rotate_mode == self.MERGE_ROTATE_LANDSCAPE and sh > sw:
                rotation = 90
                sw, sh = sh, sw

            if self.merge_size_mode == self.MERGE_SIZE_KEEP:
                out.insert_pdf(src_doc, from_page=idx, to_page=idx)
                if rotation:
                    out[out.page_count - 1].set_rotation(rotation)
                continue

            page = out.new_page(width=target_w, height=target_h)
            content_w = target_w - (CONTENT_MARGIN_LR * 2)
            content_h = target_h - (CONTENT_MARGIN_TOP + CONTENT_MARGIN_BOTTOM)
            scale = min(content_w / sw, content_h / sh) if sw and sh else 1.0
            # 확대만 방지(축소는 허용)
            if meta.source_type in ("pdf", "txt") and scale > 1.0:
                scale = 1.0
            dw, dh = sw * scale, sh * scale
            x0 = CONTENT_MARGIN_LR + (content_w - dw) / 2.0
            y0 = CONTENT_MARGIN_TOP if self.merge_align_mode == self.MERGE_ALIGN_TOP else CONTENT_MARGIN_TOP + (content_h - dh) / 2.0
            dest = fitz.Rect(x0, y0, x0 + dw, y0 + dh)
            page.show_pdf_page(dest, src_doc, idx, rotate=rotation, keep_proportion=True)

        # =========================
        # ✅ 병합 결과 페이지 번호 추가 (홀/짝 위치 분리)
        # - 현재/전체 -
        # ✅ insert_textbox 사용(정렬 안정)
        # =========================
        total_pages = out.page_count
        for i in range(total_pages):
            page = out[i]
            rect = page.rect
            text = f"- {i + 1}/{total_pages} -"

            y0 = rect.height - 40
            y1 = rect.height - 18
            margin = 50
            box_w = 200

            if (i + 1) % 2 == 1:
                # 홀수: 오른쪽
                r = fitz.Rect(rect.width - margin - box_w, y0, rect.width - margin, y1)
                align = fitz.TEXT_ALIGN_RIGHT
            else:
                # 짝수: 왼쪽
                r = fitz.Rect(margin, y0, margin + box_w, y1)
                align = fitz.TEXT_ALIGN_LEFT

            try:
                page.insert_textbox(
                    r,
                    text,
                    fontsize=9,
                    fontname="helv",
                    color=(0.45, 0.45, 0.45),
                    align=align,
                )
            except Exception:
                # 안전 폴백(중앙)
                try:
                    r2 = fitz.Rect(rect.width / 2 - box_w / 2, y0, rect.width / 2 + box_w / 2, y1)
                    page.insert_textbox(
                        r2,
                        text,
                        fontsize=9,
                        fontname="helv",
                        color=(0.45, 0.45, 0.45),
                        align=fitz.TEXT_ALIGN_CENTER,
                    )
                except Exception:
                    pass

        self._save_pdf(out, save_path)
        out.close()
        self.set_status("PDF 병합 완료")

    def action_extract(self):
        if self.work_list.count() == 0:
            return
        default_file = self.default_save_path("추출페이지", "pdf")
        p, _ = QFileDialog.getSaveFileName(self, "PDF 개별 추출 (파일명_숫자로 자동생성)", default_file, "PDF (*.pdf)")
        if not p:
            return
        self.update_last_path(p)
        base = os.path.splitext(p)[0]
        for i in range(self.work_list.count()):
            d, idx = self.work_list.item(i).data(Qt.UserRole)
            nd = fitz.open()
            try:
                nd.insert_pdf(self.docs[d].doc, from_page=idx, to_page=idx)
                self._save_pdf(nd, f"{base}_{i + 1}.pdf")
            finally:
                nd.close()
        self.set_status("개별 페이지 추출 완료")

    # ------------------------------------------------------------
    # Convert
    # ------------------------------------------------------------
    def convert_to_word(self):
        if self.work_list.count() == 0:
            return
        default_file = self.default_save_path("변환문서_결과", "docx")
        p, _ = QFileDialog.getSaveFileName(self, "Word 저장", default_file, "Word (*.docx)")
        if not p:
            return
        self.update_last_path(p)
        from docx.shared import Inches
        docx_doc = Document()

        self.progress_bar.setVisible(True)
        self.progress_bar.setMaximum(self.work_list.count())
        try:
            for i in range(self.work_list.count()):
                self.progress_bar.setFormat(f"{i+1}/{self.work_list.count()} Word 변환 중...")
                d, idx = self.work_list.item(i).data(Qt.UserRole)
                page = self.docs[d].doc[idx]
                text_content = self.extract_text(page)

                def add_image():
                    img_ptr = self.render_page_image_bytes(page, fmt="jpg")
                    docx_doc.add_picture(img_ptr, width=Inches(6.2))

                def add_text():
                    if text_content:
                        docx_doc.add_paragraph(text_content)

                if self.convert_mode == self.CONVERT_MODE_IMAGE:
                    add_image()
                elif self.convert_mode == self.CONVERT_MODE_MIXED:
                    add_image()
                    if text_content:
                        docx_doc.add_paragraph("\n[추출 텍스트]\n")
                        add_text()
                else:
                    if text_content:
                        add_text()
                    else:
                        add_image()

                if i < self.work_list.count() - 1:
                    docx_doc.add_page_break()

                self.progress_bar.setValue(i + 1)
                QApplication.processEvents()

            docx_doc.save(p)
            self.set_status("Word 변환 완료")
        finally:
            self.progress_bar.setVisible(False)

    def convert_to_ppt(self):
        if self.work_list.count() == 0:
            return
        default_file = self.default_save_path("변환발표자료_결과", "pptx")
        p, _ = QFileDialog.getSaveFileName(self, "PPT 저장", default_file, "PPT (*.pptx)")
        if not p:
            return
        self.update_last_path(p)
        from pptx.util import Pt
        from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN

        prs = Presentation()
        blank_layout = next((l for l in prs.slide_layouts if len(l.placeholders) == 0), prs.slide_layouts[6])
        d0, idx0 = self.work_list.item(0).data(Qt.UserRole)
        first_page = self.docs[d0].doc[idx0]
        prs.slide_width = Pt(first_page.rect.width)
        prs.slide_height = Pt(first_page.rect.height)
        slide_w = int(prs.slide_width)
        slide_h = int(prs.slide_height)

        self.progress_bar.setVisible(True)
        self.progress_bar.setMaximum(self.work_list.count())
        try:
            for i in range(self.work_list.count()):
                self.progress_bar.setFormat(f"{i+1}/{self.work_list.count()} PPT 변환 중...")
                d, idx = self.work_list.item(i).data(Qt.UserRole)
                page = self.docs[d].doc[idx]
                slide = prs.slides.add_slide(blank_layout)
                text_content = self.extract_text(page)

                def add_image_letterbox():
                    img_data = self.render_page_image_bytes(page, fmt="png")
                    pr = page.rect
                    img_ratio = pr.width / pr.height if pr.height else 1.0
                    slide_ratio = slide_w / slide_h if slide_h else 1.0
                    if img_ratio > slide_ratio:
                        w = slide_w
                        h = int(slide_w / img_ratio)
                    else:
                        h = slide_h
                        w = int(slide_h * img_ratio)
                    x = int((slide_w - w) / 2)
                    y = int((slide_h - h) / 2)
                    slide.shapes.add_picture(img_data, x, y, width=w, height=h)

                def add_text_box():
                    if not text_content:
                        return
                    textbox = slide.shapes.add_textbox(0, 0, slide_w, slide_h)
                    tf = textbox.text_frame
                    tf.clear()
                    tf.word_wrap = True
                    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                    tf.margin_left = Pt(24)
                    tf.margin_right = Pt(24)
                    tf.margin_top = Pt(18)
                    tf.margin_bottom = Pt(18)
                    lines = text_content.splitlines() if "\n" in text_content else [text_content]
                    first = True
                    for line in lines:
                        pgh = tf.paragraphs[0] if first else tf.add_paragraph()
                        first = False
                        pgh.text = line
                        pgh.alignment = PP_ALIGN.LEFT
                        if pgh.runs:
                            r = pgh.runs[0]
                            r.font.name = "Malgun Gothic"
                            r.font.size = Pt(20)

                if self.convert_mode == self.CONVERT_MODE_IMAGE:
                    add_image_letterbox()
                elif self.convert_mode == self.CONVERT_MODE_MIXED:
                    add_image_letterbox()
                    add_text_box()
                else:
                    if text_content:
                        add_text_box()
                    else:
                        add_image_letterbox()

                self.progress_bar.setValue(i + 1)
                QApplication.processEvents()

            prs.save(p)
            self.set_status("PPT 변환 완료")
        finally:
            self.progress_bar.setVisible(False)

    def convert_to_excel(self):
        if self.work_list.count() == 0:
            return
        default_file = self.default_save_path("페이지별_데이터_변환", "xlsx")
        p, _ = QFileDialog.getSaveFileName(self, "Excel 저장", default_file, "Excel (*.xlsx)")
        if not p:
            return
        self.update_last_path(p)
        wb = Workbook()
        wb.remove(wb.active)
        temp_paths: List[str] = []

        self.progress_bar.setVisible(True)
        self.progress_bar.setMaximum(self.work_list.count())
        try:
            for i in range(self.work_list.count()):
                self.progress_bar.setFormat(f"{i+1}/{self.work_list.count()} Excel 변환 중...")
                d, idx = self.work_list.item(i).data(Qt.UserRole)
                page = self.docs[d].doc[idx]
                ws = wb.create_sheet(title=f"Page_{i + 1}")
                text_content = self.extract_text(page)

                def put_text(start_row: int = 1):
                    if not text_content:
                        return
                    for r_idx, line in enumerate(text_content.splitlines(), start_row):
                        ws.cell(row=r_idx, column=1).value = line.strip()

                def put_image() -> int:
                    zoom = self.current_dpi / 72.0
                    matrix = fitz.Matrix(zoom, zoom).prerotate(page.rotation)
                    pix = page.get_pixmap(matrix=matrix, alpha=False)
                    try:
                        if pix.n != 3:
                            pix = fitz.Pixmap(fitz.csRGB, pix)
                    except Exception:
                        pass
                    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
                    tmp.write(pix.tobytes("png"))
                    tmp.close()
                    temp_paths.append(tmp.name)
                    ws.add_image(XLImage(tmp.name), "A1")
                    return pix.height

                if self.convert_mode == self.CONVERT_MODE_IMAGE:
                    put_image()
                elif self.convert_mode == self.CONVERT_MODE_MIXED:
                    h = put_image()
                    start_row = max(2, int(h / 20) + 2)
                    if text_content:
                        ws.cell(row=start_row - 1, column=1).value = "[추출 텍스트]"
                        put_text(start_row=start_row)
                else:
                    if text_content:
                        put_text(1)
                    else:
                        put_image()

                self.progress_bar.setValue(i + 1)
                QApplication.processEvents()

            wb.save(p)
            self.set_status("Excel 변환 완료")
        finally:
            self.progress_bar.setVisible(False)
            for tp in temp_paths:
                try:
                    os.remove(tp)
                except Exception:
                    pass

    def convert_to_text(self):
        if self.work_list.count() == 0:
            return
        default_file = self.default_save_path("텍스트추출", "txt")
        p, _ = QFileDialog.getSaveFileName(self, "Text 저장", default_file, "Text (*.txt)")
        if not p:
            return
        self.update_last_path(p)
        with open(p, "w", encoding="utf-8") as f:
            for i in range(self.work_list.count()):
                d, idx = self.work_list.item(i).data(Qt.UserRole)
                page = self.docs[d].doc[idx]
                f.write(f"--- Page {i + 1} ---\n{self.extract_text(page)}\n\n")
        self.set_status("Text 변환 완료")

    def convert_to_images(self):
        if self.work_list.count() == 0:
            return
        default_name = os.path.join(self.last_path, f"Images_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.png")
        save_path, selected_filter = QFileDialog.getSaveFileName(
            self,
            "이미지 저장 (형식 선택)",
            default_name,
            "PNG (*.png);;JPEG (*.jpg *.jpeg);;BMP (*.bmp);;TIFF (*.tif *.tiff)"
        )
        if not save_path:
            return
        self.update_last_path(save_path)

        sf = (selected_filter or "").lower()
        if "jpg" in sf or "jpeg" in sf:
            ext = "jpg"
        elif "bmp" in sf:
            ext = "bmp"
        elif "tif" in sf:
            ext = "tif"
        else:
            ext = "png"

        base_dir = os.path.dirname(save_path)
        base_name = os.path.splitext(os.path.basename(save_path))[0]

        self.progress_bar.setValue(0)
        self.progress_bar.setMaximum(self.work_list.count())
        self.progress_bar.setVisible(True)

        zoom = self.current_dpi / 72.0
        for i in range(self.work_list.count()):
            self.set_status(f"이미지 생성 중... ({i+1}/{self.work_list.count()})")
            d, idx = self.work_list.item(i).data(Qt.UserRole)
            page = self.docs[d].doc[idx]
            matrix = fitz.Matrix(zoom, zoom).prerotate(page.rotation)
            pix = page.get_pixmap(matrix=matrix, annots=True)
            try:
                pix.set_dpi(self.current_dpi, self.current_dpi)
            except Exception:
                pass
            filename = f"{base_name}_{i+1:03d}.{ext}"
            full_path = os.path.join(base_dir, filename)
            if ext == "jpg":
                pix.save(full_path, jpg_quality=85)
            else:
                pix.save(full_path)
            self.progress_bar.setValue(i + 1)
            QApplication.processEvents()

        self.progress_bar.setVisible(False)
        self.set_status(f"이미지 저장 완료 ({ext.upper()}, {self.current_dpi} DPI)")

    # ------------------------------------------------------------
    # Settings load/save
    # ------------------------------------------------------------
    def load_settings(self):
        geo = self.settings.value("geometry")
        if geo:
            self.restoreGeometry(geo)
        sp = self.settings.value("splitter_state")
        if sp:
            try:
                self.content_splitter.restoreState(sp)
            except Exception:
                pass
        else:
            self.content_splitter.setSizes([350, 350, 700])
        self.change_ui_font(self.current_font_size)
        self._update_list_icon_size()

    def closeEvent(self, e):
        for sd in self.docs:
            sd.close()
            sd.cleanup()

        # 전역 PDF 캐시 파일 정리
        for rec in list(self._pdf_render_cache.values()):
            p = rec.get("path")
            try:
                if p and os.path.exists(p):
                    os.remove(p)
            except Exception:
                pass
        self._pdf_render_cache.clear()

        self.settings.setValue("geometry", self.saveGeometry())
        self.settings.setValue("splitter_state", self.content_splitter.saveState())
        self.settings.setValue("last_path", self.last_path)
        self.settings.setValue("font_size", self.current_font_size)
        self.settings.setValue("dpi", self.current_dpi)
        self.settings.setValue("text_import_font_size", self.text_import_font_size)
        self.settings.setValue("text_mode", self.text_mode)
        self.settings.setValue("convert_mode", self.convert_mode)
        self.settings.setValue("merge_size_mode", self.merge_size_mode)
        self.settings.setValue("merge_align_mode", self.merge_align_mode)
        self.settings.setValue("merge_rotate_mode", self.merge_rotate_mode)
        self.settings.setValue("pdf_save_mode", self.pdf_save_mode)
        e.accept()


if __name__ == "__main__":
    # ✅ PyInstaller --windowed 환경에서 stdout/stderr가 None일 수 있어 방지
    # (GUI 앱 + PyMuPDF 조합에서 보고된 안정화 패턴) [1](https://github.com/pymupdf/PyMuPDF/discussions/3467)
    if getattr(sys, "stdout", None) is None:
        sys.stdout = open(os.devnull, "w")
    if getattr(sys, "stderr", None) is None:
        sys.stderr = open(os.devnull, "w")

    app = QApplication(sys.argv)
    app.setStyle("Fusion")

    # ✅ onefile에서 Qt 플러그인 탐색 경로 보험(Windows 배포 안정)
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        base = sys._MEIPASS
        QCoreApplication.addLibraryPath(os.path.join(base, "PySide6", "plugins"))

    win = IntegratedPDFMaster()
    win.show()
    sys.exit(app.exec())